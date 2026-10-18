"""문장 포함률 (#231 결정 2 = a) — docling과 **독립된** 추출기로 원문 단위를 뽑아 실제 색인 청크(OpenSearch)와 대조한다.

기준선을 docling으로 잡으면 docling이 빠뜨린 문장은 원문에서도 빠져 늘 '다 있음'이 나온다. #143이 같은 방식(pdfplumber 문장 1,866개
↔ docling 청크)으로 소항목 제목 유실을 잡았다.
  - PDF: pdfplumber. 표 영역(find_tables) 밖 글자는 줄 → 문장, 표 안은 칸 글자를 줄 → 문장(칸 하나가 쪽 전체인 경우가 있다).
  - DOCX: python-docx. 문단 → 문장, 표 칸(중첩 포함) 하나가 단위. 머리글·바닥글은 docling도 버리는 층이라 넣지 않는다.
  - XLSX: openpyxl. 값이 있는 칸 하나가 단위(병합은 왼쪽 위 칸에만 값이 있다).
그림 속 글자는 이 추출기들이 못 뽑는다 — pictures·judge(#247)가 맡는다.

분류(classify): 공백·표 구분자를 지운 문자열(pictures.norm)로 부분 일치를 본다.
  본문 — 청크 본문의 글 줄에 있다 / 표 — 청크의 표 줄(`|`로 시작)에 있다 / 헤딩 — heading_path에 있다(묶음이 헤딩을 경로로 올린다, #143)
  부분 — 글자 바이그램의 80% 이상이 한 청크에 있다(띄어쓰기·줄바꿈·하이픈 차이, 분할 경계) / 유실 후보 — 나머지.
유실 후보만 판정자(Claude)가 원문 맥락과 가장 비슷한 청크를 보고 유실/변형/추출 잡음으로 가른다.
MIN_CHARS 미만 단위(쪽 번호·'-'·단독 숫자)는 세지 않는다 — 정보가 없고 어디에나 부분 일치한다.
"""
import re
from collections import Counter
from dataclasses import asdict, dataclass
from pathlib import Path

from eval.doc_fidelity.pictures import norm as _norm

_PUNCT = re.compile(r'[\[\]()<>『』「」·•\-–—:;,.\'"‘’“”※①-⑳/→]')


def norm(text: str) -> str:
    """포함 판정용 정규화 — 공백·표 구분자에 더해 괄호·글머리표·구두점·원문자를 지운다. 공백만 지우던 첫 구현은 '(3) 임직원 …'과
    헤딩의 '(3)임직원…', '• 결과조치'와 '결과조치'를 다르게 봐 25건 중 12건을 유실 후보로 잘못 올렸다(10/18 1차 정리)."""
    return _PUNCT.sub('', _norm(text))

MIN_CHARS = 4                  # 정규화 후 글자 수 하한
PARTIAL_RATIO = 0.8            # 바이그램 겹침 하한 — 이 이상이면 '부분'(변형 존재)
_SENT_END = re.compile(r'(?<=[다요음함됨임\.\?!])\s+')


@dataclass
class Unit:
    document_id: int
    fmt: str
    element: str               # '본문' | '표'
    page: int | None
    text: str


def _sentences(line_block: str) -> list[str]:
    out = []
    for line in line_block.split('\n'):
        out += [s.strip() for s in _SENT_END.split(line) if s.strip()]
    return out


def pdf_units(document_id: int, path: Path) -> list[Unit]:
    import pdfplumber
    units: list[Unit] = []
    with pdfplumber.open(str(path)) as pdf:
        for pno, page in enumerate(pdf.pages, start=1):
            tables = page.find_tables()
            boxes = [t.bbox for t in tables]

            def outside(obj, boxes=boxes):
                x0, top, x1, bottom = obj.get('x0', 0), obj.get('top', 0), obj.get('x1', 0), obj.get('bottom', 0)
                return not any(b[0] <= x0 and x1 <= b[2] and b[1] <= top and bottom <= b[3] for b in boxes)

            text = page.filter(outside).extract_text() or ''
            units += [Unit(document_id, 'pdf', '본문', pno, s) for s in _sentences(text)]
            for t in tables:
                for row in t.extract() or []:
                    # 칸 글자도 줄·문장으로 나눈다 — 테두리 친 쪽(사내 포털 인쇄본)은 find_tables가 쪽 전체를 칸 하나짜리 표로 잡아
                    # 조문 여러 개가 한 단위가 됐고, 그런 덩어리는 어느 청크와도 안 맞아 전부 유실 후보로 떴다(급여 표준 16/30, 10/18)
                    units += [Unit(document_id, 'pdf', '표', pno, s) for c in row if c and c.strip() for s in _sentences(c)]
    return units


def docx_units(document_id: int, path: Path) -> list[Unit]:
    from docx import Document
    d = Document(str(path))
    units = [Unit(document_id, 'docx', '본문', None, s) for p in d.paragraphs for s in _sentences(p.text)]

    def cells(table):
        seen = set()
        for row in table.rows:
            for cell in row.cells:
                if id(cell._tc) in seen:          # 병합 칸은 python-docx가 같은 칸을 여러 번 준다
                    continue
                seen.add(id(cell._tc))
                own = '\n'.join(p.text for p in cell.paragraphs).strip()
                if own:
                    yield own
                for inner in cell.tables:
                    yield from cells(inner)

    for table in d.tables:
        units += [Unit(document_id, 'docx', '표', None, c.replace('\n', ' ')) for c in cells(table)]
    return units


def xlsx_units(document_id: int, path: Path) -> list[Unit]:
    import openpyxl
    wb = openpyxl.load_workbook(str(path), data_only=True)
    units = []
    for ws in wb.worksheets:
        if ws.sheet_state != 'visible':           # 숨긴 시트는 색인하지 않는다(#220 결정)
            continue
        for row in ws.iter_rows(values_only=True):
            units += [Unit(document_id, 'xlsx', '표', None, str(v).replace('\n', ' ')) for v in row
                      if v is not None and str(v).strip()]
    return units


EXTRACTORS = {'pdf': pdf_units, 'docx': docx_units, 'xlsx': xlsx_units}


def _bigrams(s: str) -> Counter:
    return Counter(s[i:i + 2] for i in range(len(s) - 1))


def best_chunk(unit_norm: str, chunks_norm: list[str]) -> tuple[int, float]:
    """바이그램 겹침이 가장 큰 청크와 그 비율(단위 바이그램 중 청크에 있는 몫)."""
    ub = _bigrams(unit_norm)
    total = sum(ub.values()) or 1
    best = (-1, 0.0)
    for i, c in enumerate(chunks_norm):
        cb = _bigrams(c)
        hit = sum(min(n, cb[g]) for g, n in ub.items()) / total
        if hit > best[1]:
            best = (i, hit)
    return best


def classify(units: list[Unit], chunks: list[dict]) -> list[dict]:
    """단위마다 분류. chunks = OpenSearch 소스(text·heading_path)."""
    body = norm('\n'.join(line for c in chunks for line in c['text'].split('\n') if not line.lstrip().startswith('|')))
    table = norm('\n'.join(line for c in chunks for line in c['text'].split('\n') if line.lstrip().startswith('|')))
    heads = norm('\n'.join(h for c in chunks for h in (c.get('heading_path') or [])))
    chunks_norm = [norm(c['text']) for c in chunks]
    out = []
    seen = set()
    for u in units:
        n = norm(u.text)
        if len(n) < MIN_CHARS or n in seen:          # 쪽마다 반복되는 머리글·바닥글(인쇄 시각·URL)은 문서당 한 번만 센다
            continue
        seen.add(n)
        if n in body:
            cls = '본문'
        elif n in table:
            cls = '표'
        elif n in heads:
            cls = '헤딩'
        else:
            idx, ratio = best_chunk(n, chunks_norm)
            cls = '부분' if ratio >= PARTIAL_RATIO else '유실 후보'
            out.append({**asdict(u), 'class': cls, 'best_chunk': idx, 'overlap': round(ratio, 3)})
            continue
        out.append({**asdict(u), 'class': cls})
    return out


LOSS_PROMPT = """아래 '원문 단위'는 문서 원본에서 뽑은 문장(또는 표 칸)이고, '색인 청크'는 검색 시스템에 들어간 글 중 가장 비슷한 것이다.
원문 단위의 정보가 색인에 들어갔는지 판정하라.
- 유실: 정보가 색인에 없다
- 변형: 띄어쓰기·줄바꿈·순서·표 형태만 다르고 정보는 있다(글자 오독이 있으면 note에 적는다)
- 잡음: 원문 추출기의 잡음(머리글·바닥글·쪽 번호·워터마크·글자 겹침)이라 색인에 없어도 정상
JSON 하나만 출력한다. 키: label(유실|변형|잡음), note

원문 단위:
{unit}

색인 청크:
{chunk}
"""
