"""문서 파일 → 청크(ChunkData) 목록. 형식별 섹션 파서 → sections.pack_sections(묶기) → 크기 분할.

형식 분기는 chunk_file 한 곳뿐이다(호출부는 확장자를 보지 않는다).
- PDF·DOCX·MD: docling 변환(상주 자식 프로세스) → 헤딩 스택으로 섹션 조립(sections.build_sections).
  docling이 실패하면 색인도 실패한다 — 저품질 경로로 조용히 바꾸지 않는다.
- TXT: 평문 하나가 섹션 하나(heading_path 없음).
- XLSX: rag/xlsx_chunking에 위임(행 묶음 + 시트 완성본, 크기 초과면 업로드 거절).
채팅 첨부(extract_text)는 사용자 대기 경로라 docling을 안 태우고 pdfplumber·python-docx·openpyxl로 평문만 뽑는다.

표는 행 단위로 나누고 묶음마다 헤더를 반복한다. 묶음 둘 이상인 표의 청크 meta에는
{"is_table", "table_key", "table_text"(표 앞 설명 + 표 + 아래 주석)}를 싣는다 — 근거 준비(rag/retriever)가 걸린 묶음을
표 전체로 바꾼다. 표 하나가 TABLE_MAX_TOKENS를 넘으면 문서를 색인하지 않는다(TableTooLargeError).
LlamaIndex 의존은 이 모듈 안에 격리한다.
"""

import logging
import os
import pickle
import re
import socket
import subprocess
import sys
import threading
from pathlib import Path
from typing import NamedTuple

import pdfplumber
from docx import Document as DocxDocument
from docx.table import Table as DocxTable
from docx.text.paragraph import Paragraph as DocxParagraph
from llama_index.core.node_parser import SentenceSplitter

from config import settings
from rag.chunking import plain, sections, split
from rag.chunking.errors import CaptionFailedError, LlmReadFailedError, PictureLimitError, TableTooLargeError
from rag.chunking.plain import pdf_image_area_ratio
from rag.chunking.split import TABLE_NOTE
from rag.chunking.tables import (
    TABLE_BATCH_CHARS,
    TABLE_BATCH_TOKENS,
    TABLE_MAX_TOKENS,
    batch_table_rows,
    to_markdown_table,
)
from rag.chunking.types import ChunkData, Item, Section, TableBlock, blocks_of, join_blocks
from rag.identifiers import HeadingSplit, is_heading_line, split_heading_line

logger = logging.getLogger(__name__)

# 밖(ingestion·routers·xlsx_chunking·eval)이 쓰는 공개 표면 — 하위 모듈로 옮겨도 여기서 이름을 유지한다.
__all__ = [
    'TABLE_BATCH_CHARS',
    'TABLE_BATCH_TOKENS',
    'TABLE_MAX_TOKENS',
    'TABLE_NOTE',
    'CaptionFailedError',
    'ChunkData',
    'Item',
    'LlmReadFailedError',
    'PictureLimitError',
    'Section',
    'TableBlock',
    'TableTooLargeError',
    'abort_conversion',
    'batch_table_rows',
    'blocks_of',
    'chunk_file',
    'count_picture_placeholders',
    'docling_block_text',
    'docling_convert',
    'docling_grid_markdown',
    'docling_grid_rows',
    'docling_picture_text',
    'extract_text',
    'join_blocks',
    'pdf_image_area_ratio',
    'to_markdown_table',
]

_DOCX_HEADING_RE = re.compile(r'^Heading (\d+)$')  # python-docx는 빌트인 헤딩을 영문명으로 준다
# Triton OCR 상수 — 서버(serving/triton-ocr/model_repository/ocr/1/model.py)와 짝이라 설정으로 빼지 않는다.
_OCR_LANG = 'korean'  # docling 기본은 ['chinese'] — 한글 표를 틀린 값으로 채운다
_OCR_SCALE = 3.0  # 크롭 렌더 배수(72dpi 기준). KserveV2OcrOptions 기본은 2.0이라 인프로세스 RapidOCR(3.0)과 맞추려면 명시해야 한다

# 캡션이 없는 그림의 자리표시 — 어떤 그림이 남는지는 count_picture_placeholders.
_PICTURE_PLACEHOLDER = '<!-- image -->'

# 그림 캡션 프롬프트 — 모든 그림에 하나. 유형별 지시는 '…라면' 조건부로만 둔다(형태를 전제한 문구는 다른 유형에 해롭다).
# 박스 아래 설명을 박스에 귀속시키라는 지시는 넣지 않는다 — 귀속을 한 칸씩 밀어 확신 있는 오답을 낸다.
_PICTURE_CAPTION_PROMPT = (
    '이미지의 내용을 한국어로 옮겨라.\n'
    '- 이미지 안의 글자를 하나도 빠뜨리지 말고 원문 그대로 옮긴다. 고쳐 쓰거나 요약하지 않는다.\n'
    '- 요소들 사이에 관계가 있으면 그 관계를 함께 적는다.\n'
    '- 관계가 분명하지 않으면 보이는 위치대로 나열만 한다. 관계가 없다는 말은 쓰지 않는다.\n'
    '- 흐름도·절차도라면 화살표 하나를 한 줄로 "<출발 박스 이름> → <도착 박스 이름>" 형식으로 적고, 화살표 위 글자(예/아니오 등)가 '
    '있으면 줄 끝에 괄호로 붙인다.\n'
    '- 조직도·계층도라면 위에서 아래로, 연결선 하나를 한 줄로 "<상위 이름> → <하위 이름>" 형식으로 적는다.\n'
    '- 차트라면 제목·축 이름·단위·범례를 먼저 적고, 값은 "가로축 항목 (계열): 값" 형식으로 항목마다 한 줄씩 적는다.\n'
    '- 이미지에 없는 내용은 쓰지 않는다.\n'
    '- 값이 비어 있는 칸·막대·점은 지어내지 말고 "값 없음"으로 적는다.\n'
    '- 머리말·맺음말·질문·제안 없이, 첫 줄부터 이미지 안의 글자로 시작한다.'
)
# 머리말·맺음말 금지는 금지할 문구를 인용하지 않고 적는다 — 인용하면 오히려 그 문구를 유도한다.


_DOCLING_LOCK = threading.Lock()
_docling_converter = None


class _PictureGuard:
    """docling 캡션 모델을 감싸 그림 장수 상한과 캡션 실패를 건다 — docling에는 둘 다 옵션이 없다.

    파이프라인 enrichment_pipe의 캡션 모델을 이 대리 객체로 바꿔 끼우고(_guard_picture_models), 나머지 속성은 원본에 넘긴다.
    - 장수: 문서의 첫 배치에서 캡션 대상(면적 하한을 넘는 그림)을 한 번 센다 — 넘으면 캡션을 하나도 붙이기 전에 끝난다.
    - 실패: 빈 캡션이면 멈춘다. 동시 N이면 그 배치에서 이미 보낸 요청은 끝까지 기다린다.
    문서 경계는 _convert_with가 arm()으로 알린다 — id(doc)로 가르면 해제된 문서의 id가 재사용될 수 있다.
    """

    def __init__(self, inner):
        self._inner = inner
        self._counted = False

    def __getattr__(self, name):
        return getattr(self._inner, name)

    def arm(self) -> None:
        self._counted = False

    def __call__(self, doc, element_batch):
        if not self._counted:
            self._counted = True
            n = _captionable_count(doc, self._inner.options.picture_area_threshold)
            if n > settings.docling_max_pictures:
                raise PictureLimitError(
                    f'그림 {n}장 — 문서당 {settings.docling_max_pictures}장을 넘는 문서는 받지 않는다'
                )
        for item in self._inner(doc, element_batch):
            desc = item.meta.description if item.meta is not None else None
            if desc is not None and not (desc.text or '').strip():
                raise CaptionFailedError('그림 캡션 실패(VLM 응답 없음·오류) — 나중에 다시 시도한다')
            yield item


def _captionable_count(doc, area_threshold: float) -> int:
    """docling이 캡션을 요청할 그림 수 — PictureDescriptionBaseModel.__call__의 면적 판정을 그대로 따른다.

    순회도 docling과 같은 doc.iterate_items()다(doc.pictures는 머리글 등 furniture까지 담아 더 크다).
    """
    from docling_core.types.doc import PictureItem  # 지연 import — docling은 톱레벨에 안 들인다

    n = 0
    for item, _ in doc.iterate_items():
        if not isinstance(item, PictureItem):
            continue
        if item.prov:
            prov = item.prov[0]
            page = doc.pages.get(prov.page_no)
            if page is not None:
                page_area = page.size.width * page.size.height
                if page_area > 0 and prov.bbox.area() / page_area < area_threshold:
                    continue
        n += 1
    return n


_XLSX_PICTURE_AREA_THRESHOLD = 0.0  # 엑셀 그림 캡션 면적 하한 — 근거는 _docling_runtime의 xlsx_opts 주석

_picture_guards: list[_PictureGuard] = []


def _guard_picture_models(converter) -> None:
    """세 형식의 파이프라인을 미리 띄우고 캡션 모델을 _PictureGuard로 바꿔 끼운다.

    initialized_pipelines는 DocumentConverter의 파이프라인 캐시다. XLSX는 캡션 면적 하한이 달라 DOCX와 파이프라인을
    공유하지 않으므로 형식마다 따로 감싼다. docling 업그레이드로 이 구조가 바뀌면 캡션 실패·장수 테스트가 먼저 깨진다.
    """
    from docling.datamodel.base_models import InputFormat
    from docling.models.picture_description_base_model import PictureDescriptionBaseModel

    for fmt in (InputFormat.PDF, InputFormat.DOCX, InputFormat.XLSX):
        converter.initialize_pipeline(fmt)
    for pipeline in converter.initialized_pipelines.values():
        for i, model in enumerate(pipeline.enrichment_pipe):
            if isinstance(model, PictureDescriptionBaseModel):
                guard = _PictureGuard(model)
                pipeline.enrichment_pipe[i] = guard
                _picture_guards.append(guard)


def _caption_headers() -> dict[str, str]:
    """캡션·OCR 요청에 붙일 HTTP 헤더 — 키가 있을 때만 Bearer(사내 VLM은 키가 없다)."""
    key = settings.vlm_caption_api_key
    return {'Authorization': f'Bearer {key}'} if key else {}


def _docling_runtime():
    """프로세스당 1회: PDF·DOCX·XLSX 변환기(모델 로딩). 변환 프로세스(_ConvertProcess) 안에서 불린다.

    메모리: 모델이 상주하고 glibc가 해제분을 바로 반납하지 않아, 여러 문서를 돈 정상상태 RSS가 단일 문서 피크보다 높다 —
    컨테이너 메모리 상한(docker-compose.yml)은 단일 문서 피크로 정하면 안 된다.
    """
    global _docling_converter
    with _DOCLING_LOCK:
        if _docling_converter is None:
            from docling.datamodel.accelerator_options import AcceleratorDevice, AcceleratorOptions
            from docling.datamodel.base_models import InputFormat
            from docling.datamodel.pipeline_options import (
                ConvertPipelineOptions,
                KserveV2OcrOptions,
                PdfPipelineOptions,
                PictureDescriptionApiOptions,
                TableFormerMode,
            )
            from docling.document_converter import (
                DocumentConverter,
                ExcelFormatOption,
                PdfFormatOption,
                WordFormatOption,
            )

            opts = PdfPipelineOptions(artifacts_path=settings.docling_artifacts_path)
            # 메모리 노브 — config.py 주석 참조. perf.page_batch_size는 쓰지 않는다: 기본 파이프라인이 그 값을 읽지 않는다.
            opts.layout_batch_size = settings.docling_layout_batch_size
            opts.queue_max_size = settings.docling_queue_max_size
            # 상용 OCR(OCR_PROVIDER=llm)이면 Triton을 안 부른다 — 사진 표·스캔 쪽은 '상용 OCR' 절이 LLM으로 읽는다.
            opts.do_ocr = settings.docling_do_ocr and settings.ocr_provider == 'triton'
            if opts.do_ocr:
                # 엔진은 원격 Triton. 기본 OcrMode는 비트맵 영역만 OCR하고 텍스트 레이어는 건드리지 않는다.
                # 요청 실패는 docling이 삼키고 result.errors에만 남기므로 _convert_with가 errors를 보고 실패로 올린다.
                opts.ocr_options = KserveV2OcrOptions(
                    url=settings.ocr_kserve_url,
                    model_name=settings.ocr_kserve_model_name,
                    transport='http',
                    lang=[_OCR_LANG],
                    scale=_OCR_SCALE,
                    timeout=settings.ocr_kserve_timeout_seconds,
                )
            opts.do_table_structure = True
            opts.table_structure_options.mode = (
                TableFormerMode.FAST if settings.docling_table_mode == 'fast' else TableFormerMode.ACCURATE
            )
            opts.table_structure_options.do_cell_matching = settings.docling_do_cell_matching
            # opts.document_timeout은 걸지 않는다 — 문서 마감은 _ConvertProcess가 프로세스째 끊는다. docling의 것은
            # PDF 쪽 루프에서만 검사하고 PARTIAL_SUCCESS(→ 결정적 실패)로 끝나 "마감 초과 = 일시 실패"와 어긋난다.
            opts.accelerator_options = AcceleratorOptions(
                num_threads=settings.docling_num_threads, device=AcceleratorDevice(settings.docling_device)
            )
            # enable_remote_services를 안 켜면 do_picture_description이 설정 단계에서 OperationNotAllowed로 죽는다.
            caption_opts = PictureDescriptionApiOptions(
                url=settings.vlm_caption_url,
                params={
                    'model': settings.vlm_caption_model,
                    'temperature': settings.vlm_caption_temperature,  # params는 요청 body에 그대로 들어간다
                    # gpt-5 계열은 'max_tokens'를 400으로 거부한다 — 사내 vLLM·Gemini는 이 이름도 상한으로 받는다
                    'max_completion_tokens': settings.vlm_caption_max_tokens,
                },
                prompt=_PICTURE_CAPTION_PROMPT,
                picture_area_threshold=settings.vlm_caption_area_threshold,
                scale=settings.vlm_caption_scale,
                headers=_caption_headers(),
                timeout=settings.vlm_caption_timeout_seconds,
                concurrency=settings.vlm_caption_concurrency,
            )
            opts.enable_remote_services = True
            opts.do_picture_description = True
            opts.picture_description_options = caption_opts
            # DOCX·XLSX는 본문이 원래 텍스트라 레이아웃·OCR 없는 SimplePipeline. 캡션 설정만 PDF와 공유한다.
            office_opts = ConvertPipelineOptions(
                enable_remote_services=True, do_picture_description=True, picture_description_options=caption_opts
            )
            # XLSX만 면적 하한 0 — 엑셀은 쪽이 없어 docling이 시트를 칸 단위로 재고 그림을 크기와 무관하게 1×1칸으로 잡아,
            # 비율이 그림 크기가 아니라 시트 데이터 양을 잰다(큰 차트도 하한에 걸려 빠진다). 장수 상한도 이 하한으로 센다.
            # 네이티브 차트 객체는 그림으로 세지지만 이미지가 없어 캡션되지 않고 자리표시만 남는다(값은 시트 칸에 있다).
            xlsx_opts = ConvertPipelineOptions(
                enable_remote_services=True,
                do_picture_description=True,
                picture_description_options=caption_opts.model_copy(
                    update={'picture_area_threshold': _XLSX_PICTURE_AREA_THRESHOLD}
                ),
            )
            _docling_converter = DocumentConverter(
                format_options={
                    InputFormat.PDF: PdfFormatOption(pipeline_options=opts),
                    InputFormat.DOCX: WordFormatOption(pipeline_options=office_opts),
                    InputFormat.XLSX: ExcelFormatOption(pipeline_options=xlsx_opts),
                }
            )
            _guard_picture_models(_docling_converter)
    return _docling_converter


# ── 상용 OCR: 사진 표·스캔 쪽을 LLM이 읽는다 ─────────────────────────────────────────
# OCR_PROVIDER=llm일 때만 탄다 — triton(기본)이면 이 절의 어느 것도 불리지 않는다.
# PDF를 쪽마다 판정해 연속된 같은 종류끼리 묶고 묶음마다 따로 변환한 뒤 합친다(DoclingDocument.concatenate —
# page_range로 변환해도 쪽 번호는 원래 번호다). 묶음 경계에서 표·헤딩이 끊길 수 있다(감수).
#   - 디지털 쪽: 기본 경로에서 OCR만 끈 것. 사진 표는 '글자 칸이 하나도 없는 표'로 나오므로 그 사진을 잘라 LLM이 칸을
#     채운다(_fill_photo_tables). 화면 캡처·흐름도는 docling이 그림으로 분류해 캡션으로 간다.
#   - 스캔 쪽: VlmPipeline — 쪽 통째로 LLM이 마크다운을 쓰고 docling이 문서로 만든다.

# 스캔 쪽 판정 — 둘 다 만족해야 스캔이다. 글자 하한이 0이 아닌 것은 스캔본에도 쪽 번호·도장 몇 글자가 글자층으로 박혀서다.
# 사진 조건은 글자 없는 표지·빈 쪽을 거른다. 보이지 않는 글자층을 깐 스캔본은 디지털로 판정돼 그 글자층이 색인된다.
_SCAN_MAX_CHARS = 20
_SCAN_MIN_IMAGE_AREA = 0.5
_OCR_LLM_SCALE = 2.0  # 쪽·사진을 LLM에 보낼 때의 해상도 배수(72dpi 기준, docling ApiVlmOptions 기본)
# docling 표 영역은 사진의 표 부분뿐이라 사진 안 제목 줄이 빠진다 → 표를 담은 사진 전체를 잘라 보낸다. 단 사진이 표보다
# 이 비율 넘게 크면(쪽 배경·레터헤드) 표 영역만 — 사진 밖 디지털 글이 두 번 색인되지 않게. 한계: CropBox가 MediaBox와 다른
# 쪽은 pdfplumber 좌표(종이 기준)와 docling 좌표(CropBox 기준)가 어긋나 표 영역만 잘린다(제목만 빠진다).
_PHOTO_MAX_AREA_RATIO = 3.0
_PHOTO_BBOX_TOL = 2.0  # pt — 표 영역이 사진 안에 드는지 볼 때의 여유(레이아웃 상자가 사진 경계를 살짝 넘는다)

# 쪽 통째 프롬프트. 절 제목에 '## '를 강제한다 — 평문 번호 제목은 docling 마크다운 해석기가 앞 번호 목록에 합쳐 번호를
# 다시 매긴다(글자가 바뀐다). 문서 제목 '#'는 맨 앞 하나 — 쪽 단위 호출이라 쪽마다 '#'가 생기지 않게.
_SCAN_PAGE_PROMPT = (
    "이 쪽을 마크다운으로 옮겨 적어라. 규칙:\n"
    "1. 문서 제목(문서 맨 앞의 큰 제목)만 '# '로 쓴다. 쪽 중간이나 둘째 쪽 이후의 제목은 문서 제목이 아니다.\n"
    "2. 절·항 제목은 번호가 있든 없든 반드시 '## '로 시작한다. 예: '## 1. 접수 원칙', '## 2.1 사진 증빙'. 번호는 원문 그대로 쓴다.\n"
    "3. 목록은 원문 표시 그대로 — 글머리표는 '- ', 번호 목록은 원문 번호 그대로.\n"
    "4. 표는 마크다운 표(| 칸 | 칸 |)로. 병합된 칸은 그 값을 해당 칸마다 반복해 적고, 빈 칸은 비워 둔다.\n"
    "5. 보이는 글자를 그대로 쓰고 없는 내용을 지어내지 마라. 설명·주석 없이 마크다운만 출력."
)
# 사진 표 프롬프트. 머리글을 지어내지 말라는 규칙은 머리글 없는 표에 '구분 | 내용' 같은 머리글을 만들어 넣는 것을 막는다.
_PHOTO_TABLE_PROMPT = (
    "이 이미지의 표를 마크다운 표(| 칸 | 칸 |)로 옮겨 적어라. 규칙:\n"
    "1. 표 밖(위·아래)에 제목·설명 같은 글 줄이 있으면 표 앞에 그 줄을 그대로 쓴다. 없으면 쓰지 않는다.\n"
    "2. 원본에 머리글 행이 있으면 첫 행에 쓰고, 없으면 머리글을 만들어 넣지 말고 첫 데이터 행을 첫 행에 쓴다.\n"
    "3. 병합된 칸은 그 값을 해당 칸마다 반복해 적고, 빈 칸은 비워 둔다.\n"
    "4. 보이는 글자를 띄어쓰기까지 그대로 쓰고 지어내지 마라. 설명 없이 출력."
)

_vlm_converter = None


class _PdfPage(NamedTuple):
    width: float
    height: float
    chars: int  # 글자층 글자 수
    images: tuple[
        tuple[float, float, float, float], ...
    ]  # 박힌 사진 (x0, top, x1, bottom) — 왼쪽 위 원점, pt, 쪽 안으로 자름


def _scan_profile(file_path: str | Path, max_pages: int | None) -> list[_PdfPage] | None:
    """쪽마다 글자 수·사진 위치(pdfplumber). 쪽 수가 max_pages를 넘으면 None — 거절은 docling이 기본 경로와 같은 실패로 한다."""
    pages = []
    with pdfplumber.open(str(file_path)) as pdf:
        if max_pages is not None and len(pdf.pages) > max_pages:
            return None
        for p in pdf.pages:
            w, h = float(p.width), float(p.height)
            boxes = (
                (
                    max(0.0, float(im['x0'])),
                    max(0.0, float(im['top'])),
                    min(w, float(im['x1'])),
                    min(h, float(im['bottom'])),
                )
                for im in p.images
            )
            pages.append(_PdfPage(w, h, len(p.chars), tuple(b for b in boxes if b[2] > b[0] and b[3] > b[1])))
    return pages


def _is_scan_page(page: _PdfPage) -> bool:
    area = page.width * page.height
    biggest = max(((x1 - x0) * (y1 - y0) for x0, y0, x1, y1 in page.images), default=0.0)
    return page.chars < _SCAN_MAX_CHARS and area > 0 and biggest / area >= _SCAN_MIN_IMAGE_AREA


def _page_runs(scan_flags: list[bool]) -> list[tuple[bool, int, int]]:
    """쪽별 스캔 여부 → 연속된 같은 종류 묶음 [(스캔인가, 첫 쪽, 끝 쪽)] — 쪽 번호는 1부터, 끝 쪽 포함(docling page_range와 같다)."""
    runs: list[tuple[bool, int, int]] = []
    for no, flag in enumerate(scan_flags, start=1):
        if runs and runs[-1][0] == flag:
            runs[-1] = (flag, runs[-1][1], no)
        else:
            runs.append((flag, no, no))
    return runs


def _vlm_runtime():
    """프로세스당 1회: 스캔 쪽용 VlmPipeline 변환기. 연결은 캡션 설정을 공유한다(config.py ocr_provider 주석)."""
    global _vlm_converter
    with _DOCLING_LOCK:
        if _vlm_converter is None:
            from docling.datamodel.base_models import InputFormat
            from docling.datamodel.pipeline_options import VlmPipelineOptions
            from docling.datamodel.pipeline_options_vlm_model import ApiVlmOptions, ResponseFormat
            from docling.document_converter import DocumentConverter, PdfFormatOption
            from docling.pipeline.vlm_pipeline import VlmPipeline

            vlm_opts = ApiVlmOptions(
                url=settings.vlm_caption_url,
                headers=_caption_headers(),
                prompt=_SCAN_PAGE_PROMPT,
                response_format=ResponseFormat.MARKDOWN,
                scale=_OCR_LLM_SCALE,
                timeout=settings.ocr_llm_timeout_seconds,
                concurrency=settings.ocr_llm_concurrency,
                # 온도는 옵션 필드로 — docling ApiVlmModel이 params 위에 vlm_options.temperature를 덮어쓴다
                temperature=settings.vlm_caption_temperature,
                params={'model': settings.vlm_caption_model, 'max_completion_tokens': settings.ocr_llm_max_tokens},
            )
            _vlm_converter = DocumentConverter(
                format_options={
                    InputFormat.PDF: PdfFormatOption(
                        pipeline_cls=VlmPipeline,
                        pipeline_options=VlmPipelineOptions(enable_remote_services=True, vlm_options=vlm_opts),
                    )
                }
            )
    return _vlm_converter


def _convert_pdf_llm(file_path: str | Path, **kwargs):
    """상용 OCR의 PDF 변환 — 쪽 묶음마다 경로를 골라 변환하고 합친다(이 절 머리 주석). 실패 분류는 _convert_with·LlmReadFailedError."""
    pages = _scan_profile(file_path, kwargs.get('max_num_pages'))
    if not pages:  # 쪽 상한 초과(None)·빈 PDF — docling이 기본 경로와 같은 방식으로 거절·처리한다
        return _convert_with(_docling_runtime(), file_path, **kwargs)
    parts = []
    for scan, first, last in _page_runs([_is_scan_page(p) for p in pages]):
        span = {} if (first, last) == (1, len(pages)) else {'page_range': (first, last)}
        if scan:
            doc = _convert_with(_vlm_runtime(), file_path, **kwargs, **span)
            _require_pages(doc, first, last)
            _titles_to_section_headers(doc)
        else:
            doc = _convert_with(_docling_runtime(), file_path, **kwargs, **span)
            _fill_photo_tables(doc, file_path, pages)
        parts.append(doc)
    if len(parts) == 1:
        return parts[0]
    from docling_core.types.doc import DoclingDocument  # 지연 import — docling은 톱레벨에 안 들인다

    return DoclingDocument.concatenate(parts)


def _require_pages(doc, first: int, last: int) -> None:
    """스캔 쪽마다 글 요소가 하나는 있어야 한다 — LLM 빈 응답(docling이 오류를 빈 글자로 삼킨다)은 그 쪽이 통째로 빈다."""
    seen = {prov.page_no for item, _ in doc.iterate_items() for prov in (getattr(item, 'prov', None) or [])}
    missing = [n for n in range(first, last + 1) if n not in seen]
    if missing:
        raise LlmReadFailedError(f'스캔 쪽 {missing} LLM 응답 없음(오류·한도 초과) — 나중에 다시 시도한다')


def _titles_to_section_headers(doc) -> None:
    """title 라벨을 section_header로 바꾼다 — LLM이 '#'로 쓴 줄은 title이 되는데 PDF 섹션 조립
    (_sections_from_docling_elements)은 section_header만 헤딩으로 봐서, 그대로면 문서 제목이 본문 줄로 떨어진다."""
    from docling_core.types.doc import DocItemLabel  # 지연 import — docling은 톱레벨에 안 들인다

    for item, _ in doc.iterate_items():
        if item.label == DocItemLabel.TITLE:
            item.label = DocItemLabel.SECTION_HEADER


def _fill_photo_tables(doc, file_path: str | Path, pages: list[_PdfPage]) -> None:
    """글자 칸이 하나도 없는 표(= OCR을 끈 경로의 사진 표)를 LLM으로 채운다. 표 밖 줄(제목)은 표 바로 앞 글로 넣는다."""
    from concurrent.futures import ThreadPoolExecutor

    from docling_core.types.doc import DocItemLabel, TableItem  # 지연 import — docling은 톱레벨에 안 들인다

    targets = [
        item
        for item, _ in doc.iterate_items()
        if isinstance(item, TableItem) and item.prov and not any((c.text or '').strip() for c in item.data.table_cells)
    ]
    if not targets:
        return
    crops = _photo_crops(doc, file_path, pages, targets)
    with ThreadPoolExecutor(max_workers=max(1, settings.ocr_llm_concurrency)) as pool:
        replies = list(pool.map(_read_photo_table, crops))
    for item, reply in zip(targets, replies):
        outside, rows = _split_table_reply(reply)
        if rows:
            _fill_table_cells(item, rows)
        if outside:
            doc.insert_text(sibling=item, label=DocItemLabel.TEXT, text=outside, prov=item.prov[0], after=False)


def _photo_box(
    table: tuple[float, float, float, float], images: tuple[tuple[float, float, float, float], ...]
) -> tuple[float, float, float, float]:
    """잘라 보낼 영역 — 표를 담은 사진 중 가장 작은 것(면적 상한 안), 없으면 표 영역 그대로. 좌표는 왼쪽 위 원점 pt."""
    l, t, r, b = table
    limit = _PHOTO_MAX_AREA_RATIO * (r - l) * (b - t)
    tol = _PHOTO_BBOX_TOL
    holders = [
        box
        for box in images
        if box[0] <= l + tol
        and box[1] <= t + tol
        and box[2] >= r - tol
        and box[3] >= b - tol
        and (box[2] - box[0]) * (box[3] - box[1]) <= limit
    ]
    return min(holders, key=lambda x: (x[2] - x[0]) * (x[3] - x[1])) if holders else table


def _photo_crops(doc, file_path: str | Path, pages: list[_PdfPage], targets) -> list:
    """표마다 잘라 보낼 이미지(PIL). pypdfium2로 쪽 하나씩 렌더해 자르고 버린다 — docling generate_page_images는
    쪽 이미지를 문서 내내 들고 있어 큰 문서에서 수백 MB가 된다."""
    import pypdfium2  # 지연 import — docling 의존성이고 상용 OCR에서만 쓴다

    crops, rendered_no, rendered = [], None, None
    pdf = pypdfium2.PdfDocument(str(file_path))
    try:
        for item in targets:
            prov = item.prov[0]
            page = pages[prov.page_no - 1]
            bb = prov.bbox.to_top_left_origin(doc.pages[prov.page_no].size.height)
            box = _photo_box((bb.l, bb.t, bb.r, bb.b), page.images)
            if rendered_no != prov.page_no:
                rendered_no, rendered = prov.page_no, pdf[prov.page_no - 1].render(scale=_OCR_LLM_SCALE).to_pil()
            crops.append(rendered.crop(tuple(round(v * _OCR_LLM_SCALE) for v in box)))
    finally:
        pdf.close()
    return crops


def _read_photo_table(image) -> str:
    """사진 표 하나를 LLM에 보낸다(docling api_image_request). 빈 응답이면 일시 실패(LlmReadFailedError).

    출력이 상한에서 잘렸거나(length) 벤더가 걸렀으면 결정적 실패(RuntimeError) — 잘린 표가 행이 빠진 채 색인되지 않게,
    그리고 같은 이미지는 다시 보내도 같다. 스캔 쪽의 같은 두 경우는 _convert_with가 PARTIAL_SUCCESS로 잡아 같은 분류다.
    """
    from docling.datamodel.base_models import VlmStopReason  # 지연 import — docling은 톱레벨에 안 들인다
    from docling.utils.api_image_request import api_image_request

    reply = api_image_request(
        image=image,
        prompt=_PHOTO_TABLE_PROMPT,
        url=settings.vlm_caption_url,
        timeout=settings.ocr_llm_timeout_seconds,
        headers=_caption_headers(),
        model=settings.vlm_caption_model,
        temperature=settings.vlm_caption_temperature,
        max_completion_tokens=settings.ocr_llm_max_tokens,
    )
    text = (reply.text or '').strip()
    if not text:
        raise LlmReadFailedError('사진 표 LLM 응답 없음(오류·한도 초과) — 나중에 다시 시도한다')
    if reply.stop_reason == VlmStopReason.LENGTH:
        raise RuntimeError(
            f'사진 표 LLM 출력이 상한({settings.ocr_llm_max_tokens}토큰)에서 잘렸다 — OCR_LLM_MAX_TOKENS를 올려 다시 올린다'
        )
    if reply.stop_reason == VlmStopReason.CONTENT_FILTERED:
        raise RuntimeError('사진 표 LLM 응답이 벤더 콘텐츠 필터에 걸렸다')
    return text


_MD_SEPARATOR_CELL = re.compile(r'^:?-+:?$')


def _split_table_reply(text: str) -> tuple[str, list[list[str]]]:
    """LLM 답 → (표 밖 줄을 이은 글, 표 행 목록). '|'로 시작하는 줄이 표다 — 구분선(|---|)은 버리고 칸 수는 가장 긴 행에 맞춰 빈 칸으로
    채운다. 코드 울타리(```)는 버리고, 표 밖 줄의 마크다운 머리 표시('#', '**')는 뗀다. 표 아래 줄(출처 등)도 표 앞 글에 합친다."""
    rows, outside = [], []
    for line in text.splitlines():
        line = line.strip()
        if not line or line.startswith('```'):
            continue
        if line.startswith('|'):
            cells = [c.strip() for c in line.strip('|').split('|')]
            if not all(_MD_SEPARATOR_CELL.match(c) for c in cells):
                rows.append(cells)
        else:
            outside.append(re.sub(r'^#+\s*', '', line).strip('*').strip())
    width = max((len(r) for r in rows), default=0)
    return ' '.join(o for o in outside if o), [r + [''] * (width - len(r)) for r in rows]


def _fill_table_cells(item, rows: list[list[str]]) -> None:
    """표 요소의 칸을 행 목록으로 바꾼다 — 첫 행이 머리글(docling_grid_rows가 첫 행을 헤더로 쓴다)."""
    from docling_core.types.doc import TableCell, TableData  # 지연 import — docling은 톱레벨에 안 들인다

    cells = [
        TableCell(
            text=text,
            start_row_offset_idx=i,
            end_row_offset_idx=i + 1,
            start_col_offset_idx=j,
            end_col_offset_idx=j + 1,
            column_header=(i == 0),
        )
        for i, row in enumerate(rows)
        for j, text in enumerate(row)
    ]
    item.data = TableData(num_rows=len(rows), num_cols=len(rows[0]), table_cells=cells)


# ── docling 텍스트 정제 — 위첨자 각주 표시 ────────────────────────────────────────

# 위첨자 각주 참조 번호. docling은 글자 크기를 버리고 'PM 2)'처럼 평문에 섞어 생성기를 헷갈리게 한다. 열거("1) …")와
# 글자 모양이 같아 정규식으로는 못 가르고, 유일한 물리적 차이인 글자 크기 비율을 pdfplumber 단어 크기로 판정한다.
_FOOTNOTE_TOKEN_RE = re.compile(r'^\d{1,2}\)$')
_FOOTNOTE_SUP_RATIO = 0.8  # 같은 줄 최대 글자 크기 대비 — 이보다 작은 'N)'만 각주 표시
_FOOTNOTE_LINE_TOL = 4.0  # 같은 줄로 묶는 top 차이(pt). 위첨자는 본문보다 조금 올라간다


def _markers_from_words(words: list[tuple[str, float, float]]) -> set[str]:
    """(단어, 글자크기, top) 목록 → 위첨자 각주 표시 집합. 한 쪽 분량을 받는다.

    top이 _FOOTNOTE_LINE_TOL 안이면 같은 줄로 묶고, 줄 최대 크기의 _FOOTNOTE_SUP_RATIO 미만인 'N)' 토큰만
    각주 표시로 본다. 각주 정의 줄("2) PM 없는 경우 생략")은 줄 전체가 작은 글자라 걸리지 않는다(본문에 남긴다).
    """
    found: set[str] = set()
    line: list[tuple[str, float]] = []
    line_top: float | None = None

    def flush():
        if not line:
            return
        body = max(size for _, size in line)
        for text, size in line:
            if _FOOTNOTE_TOKEN_RE.match(text) and size < body * _FOOTNOTE_SUP_RATIO:
                found.add(text)

    for text, size, top in sorted(words, key=lambda w: (w[2], w[0])):
        if line_top is None or abs(top - line_top) > _FOOTNOTE_LINE_TOL:
            flush()
            line, line_top = [], top
        line.append((text, size))
    flush()
    return found


def _footnote_markers(file_path: str | Path) -> dict[int, set[str]]:
    """쪽 → 위첨자 각주 표시 집합. pdfplumber로 docling 변환과 별개로 한 번 더 연다.

    best-effort: 실패하면 경고만 남기고 빈 dict — 표시가 본문에 남을 뿐이라, 보조 파서 오류로 문서를 failed 처리하는 쪽이 더 나쁘다.
    """
    out: dict[int, set[str]] = {}
    try:
        with pdfplumber.open(str(file_path)) as pdf:
            for page_no, page in enumerate(pdf.pages, start=1):
                words = [
                    (w['text'], float(w.get('size') or 0.0), float(w['top']))
                    for w in page.extract_words(extra_attrs=['size'])
                ]
                found = _markers_from_words(words)
                if found:
                    out[page_no] = found
    except Exception as e:  # noqa: BLE001 — 보조 경로. 사유는 남기고 본 경로는 계속 간다
        logger.warning('각주 사이드카 실패(각주 표시 제거 없이 진행) %s: %s', file_path, e)
    return out


def _strip_footnote_markers(text: str, markers: set[str]) -> str:
    """본문에서 각주 참조 'N)'를 뗀다 — 단어·닫는 괄호 **뒤에 붙은 것만**. 줄 머리의 정의 줄("2) PM …")과
    같은 번호가 아닌 열거, 더 큰 숫자의 일부("12)")는 그대로다. 같은 쪽에 같은 번호의 열거가 단어 뒤에 붙어 있으면 그것도 지워진다(감수)."""
    for marker in sorted(markers, key=len, reverse=True):
        # 앞의 공백은 같은 줄 안에서만([ \t]) — 줄바꿈을 허용하면 다음 줄 머리의 정의 줄 "2) PM …"까지 지운다.
        # (?<!\d): "서식 12)"의 "2)"는 각주가 아니라 숫자의 일부다.
        pattern = r'(?<=[\w\)\]）」])[ \t]?(?<!\d)' + re.escape(marker) + r'(?=$|[\s,.;:)>\]）」→])'
        text = re.sub(pattern, '', text)
    return text


def _docling_norm(text: str) -> str:
    """docling 텍스트의 NBSP(\\xa0)를 보통 공백으로. 표 셀·헤딩에 섞여 들어와 검색 토큰을 깨뜨린다."""
    return re.sub(r'[  ]+', ' ', text).strip()


def docling_grid_markdown(table, doc=None) -> str | None:
    """docling 표 → markdown, 셀 원문(grid)에서 직접 그린다. docling 표 렌더의 정의점(첫 행 = 헤더, 열이 없으면 None).

    - 병합 칸은 범위의 모든 칸에 값을 둔다(가로·세로 모두) — 첫 칸에만 두면 나머지 칸이 빈칸으로 읽히거나, 청크 경계에서
      잘린 아래 행이 분류어를 잃는다. 대가는 긴 병합 셀의 중복이다.
    - export_to_dataframe을 쓰지 않는다 — column_header 표시가 없는 표에 열 이름 0,1,2…를 붙이고, Word 서식 칸을
      '**항목**'처럼 마크다운 굵게로 내보낸다(셀 원문에는 '**'가 없다).
    - 값이 하나도 없는 행은 빼고, 빈 셀은 자리를 지킨다(to_markdown_table).
    - 칸 안의 그림은 그 칸 글자 뒤에 캡션(없으면 자리표시)을 한 줄로 붙인다(doc가 있을 때) — 캡처 설명이 제 행과 붙어 있게.
      여러 줄 캡션의 줄바꿈은 사라진다(마크다운 표는 행 = 한 줄).
    """
    rows = docling_grid_rows(table, doc)
    return to_markdown_table(rows[0], rows[1:]) if rows else None


def docling_grid_rows(table, doc=None) -> list[list[str]]:
    """docling 표 → 행 목록(첫 행 = 헤더). docling_grid_markdown의 행 단계(규칙은 그 docstring) — 행 묶음은 글로
    바꾸기 전에 나눠야 해서 따로 둔다(batch_table_rows)."""
    rows = [[_grid_cell_text(cell, doc) for cell in row] for row in table.data.grid]
    return [row for row in rows if any(row)]


_XLSX_MIDNIGHT_RE = re.compile(r'^(\d{4}-\d{2}-\d{2}) 00:00:00$')


def _cell_text_norm(text: str) -> str:
    """칸 글자 정리: 줄바꿈은 공백으로(markdown 표는 행 = 한 줄), 엑셀 날짜 칸의 ` 00:00:00` 꼬리는 뗀다
    (docling이 datetime을 str로 바꾸며 붙인다). 자정이 아닌 실제 시각은 그대로 둔다."""
    text = _docling_norm(text.replace('\n', ' '))
    return _XLSX_MIDNIGHT_RE.sub(r'\1', text)


def _grid_cell_text(cell, doc) -> str:
    """표 칸 하나의 글자 — 칸 원문(_cell_text_norm) + 칸 안 그림의 설명(docling_grid_markdown 참조)."""
    from docling_core.types.doc import PictureItem, RichTableCell  # 지연 import — docling은 톱레벨에 안 들인다

    parts = [_cell_text_norm(cell.text)]
    if doc is not None and isinstance(cell, RichTableCell):
        parts += [
            re.sub(r'\s+', ' ', docling_picture_text(item))
            for item, _ in doc.iterate_items(root=cell.ref.resolve(doc))
            if isinstance(item, PictureItem)
        ]
    return ' '.join(p for p in parts if p)


def docling_picture_text(element) -> str:
    """그림 요소의 본문 한 줄 — 캡션, 없으면 자리표시(캡션 대상이 아닌 그림)."""
    return _docling_norm(_docling_picture_caption(element)) or _PICTURE_PLACEHOLDER


def docling_block_text(item, doc) -> str:
    """docling 요소 하나 → 본문 텍스트(없으면 ''). DOCX·MD·XLSX가 공유하는 요소 종류 분기의 정의점.

    목록 항목은 docling이 text에서 뗀 번호(marker)를 되붙인다. 헤딩 판정은 형식마다 달라 호출부가 한다.
    """
    from docling_core.types.doc import ListItem, PictureItem, TableItem, TextItem  # 지연 import

    match item:
        case TableItem():
            return docling_grid_markdown(item, doc) or ''
        case PictureItem():
            return docling_picture_text(item)
        case ListItem(marker=marker) if marker:
            return _docling_norm(f'{marker} {item.text}')
        case TextItem():
            return _docling_norm(item.text)
    return ''


def count_picture_placeholders(chunks) -> int:
    """청크들에 남은 그림 자리표시 수 — 캡션 대상에서 빠진 그림(면적 하한 미만·MD 그림·엑셀 차트 객체)의 수.

    캡션이 붙은 그림은 본문과 구분되지 않아 셀 수 없다. 지표 기록은 호출부(rag/ingestion.py)가 한다 — 이 모듈은
    0층이라 rag/metrics.py를 import하지 않는다.
    """
    return sum(c.text.count(_PICTURE_PLACEHOLDER) for c in chunks)


def _docling_picture_caption(element) -> str:
    """그림 요소의 VLM 캡션(없으면 ''). docling은 캡션을 element.text가 아니라 annotations에 담는다 — 첫 비어 있지 않은 것."""
    for ann in getattr(element, 'annotations', None) or []:
        text = (getattr(ann, 'text', '') or '').strip()
        if text:
            return text
    return ''


def _docling_elements(doc, footnote_markers: dict[int, set[str]] | None = None):
    """docling PDF 문서 → (label, text, page, table_rows) 흐름(읽기 순서). table_rows는 표일 때만 행 목록(docling_grid_rows).

    _sections_from_docling_elements와 나눈 것은 매핑 규칙을 docling 없이 튜플만으로 단위 테스트하기 위해서다.
    footnote_markers(쪽 → 각주 표시)는 텍스트 요소에서만 뗀다 — 표·그림은 대상이 아니다.
    """
    footnote_markers = footnote_markers or {}
    for element, _depth in doc.iterate_items():
        label = element.label.value if hasattr(element.label, 'value') else str(element.label)
        page = element.prov[0].page_no if getattr(element, 'prov', None) else None
        if label == 'table':
            yield label, '', page, docling_grid_rows(element, doc) or None
        elif label == 'picture':
            # 캡션이 없는 그림은 빈 문자열이 나가고, 소비자가 자리표시로 바꾼다.
            yield label, _docling_picture_caption(element), page, None
        else:
            text = getattr(element, 'text', '') or ''
            # docling은 목록 번호·글머리('2.1', '-', '(1)')를 text에서 떼어 marker에 둔다 — export_to_markdown처럼 다시 붙인다.
            marker = getattr(element, 'marker', '') or ''
            if label == 'list_item' and marker:
                text = f'{marker} {text}'
            if page in footnote_markers:
                text = _strip_footnote_markers(text, footnote_markers[page])
            yield label, text, page, None


def _join_split_article_titles(elements):
    """닫는 괄호만 다음 요소로 떨어진 조 제목을 잇는다 — '제 7 조 ( 징계종류' + ')' → '제 7 조 ( 징계종류 )'."""
    elements = list(elements)
    i = 0
    while i < len(elements):
        label, text, page, table = elements[i]
        nxt = elements[i + 1] if i + 1 < len(elements) else None
        t = _docling_norm(text or '')
        if (
            label in ('section_header', 'text')
            and nxt is not None
            and _docling_norm(nxt[1] or '') in (')', '）')
            and (head := split_heading_line(t))
            and head.head == t
            and not any(c in t for c in ')）')
        ):
            yield label, f'{t} {_docling_norm(nxt[1])}', page, table
            i += 2
            continue
        yield label, text, page, table
        i += 1


def _sections_from_docling_elements(elements) -> list[Section]:
    """(label, text, page, table) 흐름 → Section. PDF 헤딩 층 복원 규칙의 정의점. table은 행 목록(TableBlock으로
    운반)이거나 markdown 문자열(그대로 글 블록 — 튜플만 먹이는 테스트용).

    docling은 장·조를 구분하지 않으므로 여기서 층을 만든다:
      첫 헤딩 = 문서 제목(L1) / '제 N 장' = L2 / 그 외 헤딩 = L3 (장이 아직 없으면 L2).
    그래야 sections.pack_sections가 같은 장 아래 조항끼리 묶는다. 그림은 캡션(없으면 자리표시), 나머지 텍스트 요소는 본문 줄.

    조 제목 보정 — docling 레이아웃 모델은 본문과 같은 글꼴의 조 제목을 들쭉날쭉 section_header/text로 판정한다. 두 꼴만 고친다:
      ① 문서 제목 뒤의 text 요소가 조 제목 꼴(rag/identifiers.split_heading_line)로 시작하면 제목으로 올리고, 같은 요소에
         이어진 글은 그 조의 본문으로 둔다. 한계: 조 제목 꼴로 시작하는 본문 문단(쪽 경계에서 끊긴 문장·목차 줄)도 올라간다(글은 남는다).
      ② 닫는 괄호가 다음 요소로 떨어진 조 제목은 하나로 잇는다(_join_split_article_titles) — 안 이으면 ')'가 경로에 들어간다.
    docling이 문서 제목까지 본문으로 낸 문서에서는 보정이 걸리지 않는다. 조문 체계 밖의 제목('1. 목적')은 글자로 가를 단서가 없어 보정하지 않는다.
    """
    items: list[Item] = []
    seen_title = False
    seen_chapter = False
    for label, text, page, table_markdown in _join_split_article_titles(elements):
        if label == 'text' and seen_title and (head := split_heading_line(_docling_norm(text))):
            # ① 조 제목 꼴 본문 요소 → 제목(+ 이어진 글은 본문)
            items.append((page, head.head, 3 if seen_chapter else 2))
            if head.rest:
                items.append((page, head.rest, None))
            continue
        if label == 'section_header':
            text = _docling_norm(text)
            if not text:
                continue
            if not seen_title:
                level, seen_title = 1, True
            elif sections.CHAPTER_RE.match(text):
                level, seen_chapter = 2, True
            else:
                level = 3 if seen_chapter else 2
            items.append((page, text, level))
        elif label == 'table':
            if not table_markdown:
                continue
            if isinstance(table_markdown, str):  # markdown 문자열 — 튜플만 먹이는 단위 테스트 호환
                items.append((page, table_markdown, None))
                continue
            block = TableBlock(table_markdown, [page] * len(table_markdown))
            # 쪽을 넘긴 표 잇기 — docling은 쪽마다 표를 따로 낸다. 바로 앞 항목이 표이고 헤더가 같을 때만 뒤 표의 헤더를
            # 떼고 행을 잇는다. 행별 쪽 번호는 각자 것을 지킨다(row_pages).
            if items and isinstance(items[-1][1], TableBlock) and items[-1][1].rows[0] == block.rows[0]:
                prev_page, prev_block, _ = items[-1]
                items[-1] = (
                    prev_page,
                    TableBlock(prev_block.rows + block.rows[1:], prev_block.row_pages + block.row_pages[1:]),
                    None,
                )
                continue
            items.append((page, block, None))
        elif label == 'picture':
            caption = _docling_norm(text)
            items.append((page, caption or _PICTURE_PLACEHOLDER, None))
        else:
            text = _docling_norm(text)
            if text:
                items.append((page, text, None))
    return sections.build_sections(items)


def _convert_local(file_path: str | Path, **kwargs):
    """이 프로세스에서 docling 변환 한 번 → DoclingDocument. 변환 프로세스(_convert_child_main)가 부른다.

    상용 OCR(OCR_PROVIDER=llm)의 PDF는 _convert_pdf_llm으로, 나머지는 기본 변환기로 _convert_with를 거친다.
    kwargs는 converter.convert로 넘긴다(PDF의 max_num_pages — 쪽 개념이 없는 형식엔 주지 않는다).
    """
    if settings.ocr_provider == 'llm' and Path(file_path).suffix.lower() == '.pdf':
        return _convert_pdf_llm(file_path, **kwargs)
    return _convert_with(_docling_runtime(), file_path, **kwargs)


def _convert_with(converter, file_path: str | Path, **kwargs):
    """변환기 하나로 변환 + 불완전 결과 가드의 정의점.

    raises_on_error=True만으로는 부족하다 — docling은 PARTIAL_SUCCESS(원격 OCR 실패 등)에 예외를 안 던지므로 status와
    errors를 직접 보고 올린다. _PictureGuard의 예외는 docling이 RuntimeError로 감싸므로 원래 예외로 풀어 다시 던진다
    (감싼 채면 outbox.is_transient가 결정적으로 본다). 상용 OCR은 쪽 묶음마다 부르므로 그림 장수 상한도 묶음마다 센다.
    """
    for guard in _picture_guards:
        guard.arm()
    try:
        result = converter.convert(str(file_path), raises_on_error=True, **kwargs)
    except Exception as e:
        cause = e.__cause__
        while cause is not None:
            if isinstance(cause, (PictureLimitError, CaptionFailedError)):
                raise cause from None
            cause = cause.__cause__
        raise
    from docling.datamodel.base_models import ConversionStatus  # 지연 import — docling은 톱레벨에 안 들인다

    if result.status != ConversionStatus.SUCCESS or result.errors:
        detail = '; '.join(
            f'{e.category.value}/{e.module_name}: {(e.error_message or "")[:100]}' for e in result.errors[:3]
        )
        raise RuntimeError(f'docling 변환 불완전 — status={result.status.value}, errors={len(result.errors)}: {detail}')
    return result.document


def _picklable(exc: BaseException) -> BaseException:
    """부모로 보낼 수 있는 예외 — 피클이 안 되면 같은 문구의 RuntimeError(결정적)로 바꾼다."""
    try:
        pickle.dumps(exc)
        return exc
    except Exception:  # noqa: BLE001 — 직렬화 불가 예외는 종류를 잃는다(문구는 남긴다)
        return RuntimeError(f'{type(exc).__name__}: {exc}')


def _convert_child_main(fd: int) -> None:
    """변환 프로세스의 본체 — 요청 (경로, kwargs)를 받아 ('ok', DoclingDocument) 또는 ('err', 예외)를 돌려준다.

    첫 메시지는 부모의 설정 스냅숏(_CHILD_SETTING_PREFIXES)이다 — 자식은 config를 .env에서 새로 읽으므로, 부모가 바꾼 값
    (테스트의 monkeypatch 등)이 자식에 닿는 유일한 길이다. 모델 로딩은 'ready' 전에 끝내 첫 문서의 마감을 깎지 않는다.
    """
    from multiprocessing.connection import Connection

    conn = Connection(fd)
    for key, value in conn.recv().items():
        setattr(settings, key, value)
    try:
        _docling_runtime()
    except Exception as e:  # noqa: BLE001 — 기동 실패도 부모에 사유로 돌려준다
        conn.send(('err', _picklable(e)))
        return
    conn.send(('ready', None))
    while True:
        try:
            file_path, kwargs = conn.recv()
        except EOFError:  # 부모가 끊었다(재기동·종료) — 부모가 죽어도 여기서 빠진다
            return
        try:
            conn.send(('ok', _convert_local(file_path, **kwargs)))
        except Exception as e:  # noqa: BLE001 — 문서 단위 실패는 부모가 분류한다(outbox.is_transient)
            conn.send(('err', _picklable(e)))


_CHILD_START_TIMEOUT_SECONDS = 300  # 변환 프로세스 기동(모델 로딩) 상한 — 첫 실행의 모델 다운로드까지 품는 여유
_PROJECT_ROOT = Path(__file__).resolve().parents[2]  # rag/chunking/ 아래라 두 단계 위가 저장소 루트


# 변환 프로세스가 쓰는 설정의 접두사 — 부모는 이 설정만 자식에 넘기고, 이것이 바뀌었을 때만 자식을 다시 띄운다
# (변환과 무관한 설정을 바꿀 때마다 콜드 기동하지 않게). 이 모듈이 읽는 settings 필드가 모두 이 접두사 안인지는
# tests/test_chunking_docling.py가 검사한다 — 빠뜨리면 값을 바꿔도 자식이 옛 값으로 돈다.
_CHILD_SETTING_PREFIXES = (
    'docling_',
    'ocr_',
    'vlm_caption_',
)


def _child_settings() -> dict:
    return {k: v for k, v in settings.model_dump().items() if k.startswith(_CHILD_SETTING_PREFIXES)}


class _ConvertProcess:
    """docling 변환을 맡는 상주 자식 프로세스 하나(이 모듈을 쓰는 프로세스마다 하나).

    프로세스인 이유: docling에는 취소 API가 없어 스레드로 돌리면 마감을 넘겨도 멈출 수 없다. 프로세스는 kill하고 다음
    요청 때 새로 띄운다. 자식은 변환만 한다 — 임베딩·쓰기·커밋은 부모에서 그 뒤라, 변환 중 kill해도 반쪽 색인은 없다.
    - 마감: settings.docling_document_timeout_seconds(전 형식, 캡션 포함). 넘기면 TimeoutError(일시 실패).
    - 동시성 1: 잠금 하나로 요청을 직렬화한다.
    - 띄우는 법: `python -c` + 소켓 한 쌍 위의 multiprocessing Connection. spawn은 부모의 __main__ 스크립트를 자식에서
      다시 import하고(보호 없는 스크립트는 본문이 한 번 더 돈다), fork는 부모의 스레드·잠금 상태를 복제해 교착 위험이 있다.
    - 설정 스냅숏이 바뀌면 재기동한다 — 자식은 기동 때의 설정으로 변환기를 만든다.
    """

    def __init__(self):
        self._lock = threading.Lock()
        self._proc: subprocess.Popen | None = None
        self._conn = None
        self._snapshot: dict | None = None

    def convert(self, file_path: str | Path, kwargs: dict):
        with self._lock:
            snapshot = _child_settings()
            if self._proc is None or self._proc.poll() is not None or snapshot != self._snapshot:
                self._start(snapshot)
            deadline = settings.docling_document_timeout_seconds
            self._conn.send((str(Path(file_path).resolve()), kwargs))  # 자식의 cwd는 저장소 루트다
            if not self._conn.poll(deadline):
                self._stop()
                raise TimeoutError(f'docling 변환이 {deadline:.0f}초를 넘겨 중단했다(변환 프로세스 재기동)')
            try:
                kind, value = self._conn.recv()
            except EOFError:  # 자식이 죽었다(OOM kill 등) — poll은 EOF에도 True다
                code = self._exit_code()
                self._stop()
                raise RuntimeError(f'docling 변환 프로세스가 비정상 종료했다(exitcode={code})') from None
        if kind == 'ok':
            return value
        raise value

    def _start(self, snapshot: dict) -> None:
        from multiprocessing.connection import Connection

        self._stop()
        parent_sock, child_sock = socket.socketpair()
        env = {
            **os.environ,
            'PYTHONPATH': os.pathsep.join(p for p in (str(_PROJECT_ROOT), os.environ.get('PYTHONPATH')) if p),
        }
        try:
            self._proc = subprocess.Popen(
                [
                    sys.executable,
                    '-c',
                    'import sys; from rag.chunking import _convert_child_main; _convert_child_main(int(sys.argv[1]))',
                    str(child_sock.fileno()),
                ],
                pass_fds=(child_sock.fileno(),),
                cwd=_PROJECT_ROOT,
                env=env,
            )
        except BaseException:
            parent_sock.close()
            raise
        finally:
            child_sock.close()
        self._conn = Connection(parent_sock.detach())
        self._conn.send(snapshot)
        if not self._conn.poll(_CHILD_START_TIMEOUT_SECONDS):
            self._stop()
            raise TimeoutError(f'docling 변환 프로세스 기동이 {_CHILD_START_TIMEOUT_SECONDS}초를 넘겼다')
        try:
            kind, value = self._conn.recv()
        except EOFError:
            code = self._exit_code()
            self._stop()
            raise RuntimeError(f'docling 변환 프로세스가 기동 중 종료했다(exitcode={code})') from None
        if kind != 'ready':
            self._stop()
            raise value
        self._snapshot = snapshot

    def abort(self) -> None:
        """진행 중인 변환을 끊는다 — **잠금 없이** 부른다(잠금은 기다리는 스레드가 쥐고 있다). 다른 스레드에서 불려도
        안전하다: kill만 하고, 기다리던 convert()가 EOF로 깨어나 정리(_stop)하고 잠금을 푼다."""
        proc = self._proc
        if proc is not None and proc.poll() is None:
            proc.kill()

    def _exit_code(self):
        try:
            return self._proc.wait(5)
        except subprocess.TimeoutExpired:
            return None

    def _stop(self) -> None:
        if self._conn is not None:
            self._conn.close()
        if self._proc is not None and self._proc.poll() is None:
            self._proc.kill()  # 자식은 지킬 상태가 없다 — C 코드에 묶여 있어도 끊기는 SIGKILL로
            self._proc.wait(10)
        self._proc = self._conn = self._snapshot = None


_convert_process = _ConvertProcess()


def abort_conversion() -> None:
    """진행 중인 docling 변환을 끊는다(변환 프로세스 kill). 색인 잡이 취소될 때 rag/ingestion.py가 부른다."""
    _convert_process.abort()


def docling_convert(file_path: str | Path, **kwargs):
    """docling 변환 한 번 → DoclingDocument(변환 프로세스에 맡긴다). PDF·DOCX·MD·XLSX가 공유한다.

    실패 분류(outbox.is_transient): TimeoutError·CaptionFailedError·LlmReadFailedError는 일시,
    PictureLimitError·불완전 변환 RuntimeError는 결정적.
    rag/xlsx_chunking이 이 이름을 import한다 — 테스트가 막으려면 이 모듈과 xlsx_chunking 둘 다 패치해야 한다.
    """
    return _convert_process.convert(file_path, kwargs)


def _docling_sections(file_path: str | Path) -> list[Section]:
    """PDF → docling(레이아웃 ML + TableFormer, 그림 캡션, 표 이미지 OCR) → Section.

    인쇄 머리글·바닥글은 docling이 furniture로 분류해 iterate_items()에 나오지 않는다. 헤딩 층 복원은
    _sections_from_docling_elements가 한다. 실패는 그대로 올린다 — 저품질 폴백 없음.
    """
    document = docling_convert(file_path, max_num_pages=settings.docling_max_num_pages)
    # 각주 사이드카는 변환이 성공한 뒤에만 연다 — 변환 실패 사유가 pdfplumber 오류에 가려지지 않게,
    # 그리고 가짜 변환기 테스트가 파일 없이도 변환 가드를 검증하게.
    footnote_markers = _footnote_markers(file_path)
    return _sections_from_docling_elements(_docling_elements(document, footnote_markers))


def _docx_items(file_path: str | Path):
    """DOCX body를 문서 순서대로 (is_heading, level, text, from_table)로 흘린다(python-docx). 채팅 첨부(_docx_text) 전용.

    doc.paragraphs와 doc.tables를 따로 훑으면 표가 문서 끝으로 밀리므로 body 자식을 직접 순회한다.
    """
    doc = DocxDocument(str(file_path))
    for child in doc.element.body.iterchildren():
        tag = child.tag.split('}')[-1]
        if tag == 'p':
            para = DocxParagraph(child, doc)
            text = para.text.strip()
            if not text:
                continue
            # 스타일 정의가 빠진 docx는 style이 None이다
            m = _DOCX_HEADING_RE.match(getattr(para.style, 'name', None) or '')
            yield (bool(m), int(m.group(1)) if m else 0, text, False)
        elif tag == 'tbl':
            for row in DocxTable(child, doc).rows:
                # 빈 셀도 자리를 지킨다 — 걸러내면 뒷 열이 앞으로 당겨져 헤더와 어긋난다. 완전 빈 행만 건너뛴다.
                cells = [c.text.strip() for c in row.cells]
                if any(cells):
                    yield (False, 0, ' | '.join(cells), True)  # from_table — 폴백 헤딩 승격 대상에서 제외


def _docx_text(file_path: str | Path) -> str:
    """DOCX 전체를 평문 하나로(채팅 첨부용). 헤딩 포함, 문서 순서 보존."""
    return '\n'.join(text for _, _, text, _ in _docx_items(file_path))


_DOCX_TITLE_MAX = 60


def _split_heading_head(text: str) -> tuple[str, str]:
    """조문·장 문단 → (헤딩 머리, 나머지 본문). 조는 첫 줄의 '제N조(제목)'까지(identifiers.split_heading_line),
    장은 첫 줄 전체.

    문단 전체를 헤딩으로 올리면 같은 문단에 이어진 조문 본문이 heading_path로 들어가고, 뒤에 본문 줄이 없으면
    섹션이 생기지 않아 색인에서 사라진다.
    """
    first, _, rest_lines = text.partition('\n')
    head_split = split_heading_line(first) or HeadingSplit(first.strip(), '')  # 장('제 N 장 …')은 첫 줄 전체가 머리
    return head_split.head, (head_split.rest + ('\n' + rest_lines if rest_lines else '')).strip()


def _docx_fallback_headings(items: list[tuple[bool, int, str, bool]]) -> list[tuple[bool, int, str, bool]]:
    """Heading 스타일이 하나도 없는 docx의 조문 패턴 승격. 입력에 헤딩이 하나라도 있으면 그대로 돌려준다.

    스타일 없는 문서는 전체가 섹션 하나가 되어 분할기가 아무 데서나 자른다. 층은 PDF(_sections_from_docling_elements)와 같다:
      제목 = 첫 문단의 첫 줄(조문 패턴이 아니고 _DOCX_TITLE_MAX자 이하) → L1 — 휴리스틱이라 "목 차"도 제목이 된다
      '제 N 장' → L2 / '제 N 조' → L3(장이 아직 없으면 L2). 제목이 없으면 전부 한 단계 위.
    승격하지 않는 것(본문으로 남는다): 'N.'·'N.N' 번호 문단(목록 항목과 못 가른다), 표 행(from_table),
    인용 문단·괄호 없이 적은 조('제8조 목적') — 조 판정은 rag.identifiers.split_heading_line 한 곳이다.
    헤딩은 머리('제N조(제목)')까지만, 같은 문단에 이어진 본문은 본문 줄로 분리한다(_split_heading_head).
    """
    if any(is_heading for is_heading, _, _, _ in items):
        return items
    out: list[tuple[bool, int, str, bool]] = []
    has_title = False
    seen_chapter = False
    for i, (is_heading, level, text, from_table) in enumerate(items):
        if from_table:
            out.append((is_heading, level, text, from_table))
            continue
        first_line = text.split('\n', 1)[0].strip()
        if (
            i == 0
            and len(first_line) <= _DOCX_TITLE_MAX
            and not sections.CHAPTER_RE.match(first_line)
            and not is_heading_line(first_line)
        ):
            has_title = True
            out.append((True, 1, first_line, False))
            rest = text[len(text.split('\n', 1)[0]) :].strip()
            if rest:
                out.append((False, 0, rest, False))
            continue
        # 인용 문단("제98조 및 제99조에 따라")을 제목으로 올리면 뒤에 곧바로 다른 제목이 올 때 그 문장이 색인에서 사라진다.
        if sections.CHAPTER_RE.match(text) or is_heading_line(first_line):
            if sections.CHAPTER_RE.match(text):
                seen_chapter = True
                level = 2
            else:
                level = 3 if seen_chapter else 2
            head, rest = _split_heading_head(text)
            out.append((True, level, head, False))
            if rest:
                out.append((False, 0, rest, False))
        else:
            out.append((False, 0, text, False))
    if not has_title:
        out = [(h, (lv - 1 if h else lv), t, ft) for h, lv, t, ft in out]
    return out


class _DocxItem(NamedTuple):
    """DOCX 항목 — _docx_items·_docx_docling_items와 같은 꼴이라 조문 폴백을 공유한다(튜플로 풀어 쓰기도 통한다)."""

    is_heading: bool
    level: int  # 헤딩 레벨(본문 줄은 0)
    text: str
    from_table: bool  # 표·그림 — 폴백 헤딩 승격에서 제외하는 표지


def _docx_table_rows(table, doc) -> list[list[str]]:
    """DOCX·MD 표 → 행 목록(첫 행 헤더), 표 안의 표는 평탄화한다.

    표 안 요소는 따로 건너뛰므로(_table_descendants) 안쪽 표는 여기서 바깥 표에 실어야 사라지지 않는다. 모양은 PDF
    TableFormer와 같다: 안쪽 표가 든 열은 안쪽 열 수만큼 넓어지고 그 행은 안쪽 행 수만큼 늘며, 같은 행의 다른 칸은 세로로,
    다른 행의 그 열 칸은 넓어진 열마다 반복한다(병합 반복 규칙과 같다). 안쪽 표가 둘 이상인 칸은 위아래로 이어 붙인다.
    """
    from docling_core.types.doc import RichTableCell, TableItem  # 지연 import

    grid = table.data.grid
    nested: dict[tuple[int, int], list[list[str]]] = {}
    for r, row in enumerate(grid):
        for c, cell in enumerate(row):
            if not (
                isinstance(cell, RichTableCell) and cell.start_row_offset_idx == r and cell.start_col_offset_idx == c
            ):
                continue
            inner = [it for it, _ in doc.iterate_items(root=cell.ref.resolve(doc)) if isinstance(it, TableItem)]
            deeper = {d.self_ref for it in inner for d, _ in doc.iterate_items(root=it) if d is not it}
            inner = [it for it in inner if it.self_ref not in deeper]  # 3단 중첩은 재귀가 맡는다 — 중복 금지
            rows = [row2 for it in inner for row2 in _docx_table_rows(it, doc)]
            if rows:
                own = _cell_text_norm(cell.text)  # 칸에 안쪽 표와 함께 적힌 글자(라벨)는 첫 행으로 남긴다
                nested[(r, c)] = ([[own] + [''] * (len(rows[0]) - 1)] if own else []) + rows
    base = [[_grid_cell_text(cell, doc) for cell in row] for row in grid]
    if not nested:
        return [row for row in base if any(row)]
    width = {
        c: max(len(x) for (_, cc), rows in nested.items() if cc == c for x in rows)  # 열 c → 안쪽 열 수(최대)
        for c in {c for (_, c) in nested}
    }
    out: list[list[str]] = []
    for r, row in enumerate(base):
        inner_rows = max((len(rows) for (rr, _), rows in nested.items() if rr == r), default=0)
        for i in range(max(inner_rows, 1)):
            line: list[str] = []
            for c, text in enumerate(row):
                if c in width:
                    rows = nested.get((r, c))
                    if rows:
                        cells = rows[i] if i < len(rows) else []
                        line.extend((cells + [''] * width[c])[: width[c]])
                    else:
                        line.extend([text] * width[c])  # 넓어진 열에 걸친 칸 = 가로 병합 — 칸마다 반복
                else:
                    line.append(text)  # 안쪽 행마다 반복 — 세로 병합 반복 규칙(docling_grid_rows)과 같다
            out.append(line)
    return [row for row in out if any(row)]


def _table_descendants(doc) -> set[str]:
    """표 칸 안에 든 요소들의 self_ref. 표 자신은 빼고, 중첩 표는 바깥 표의 후손이라 들어간다(_docx_table_rows가 평탄화해 싣는다)."""
    return {item.self_ref for table in doc.tables for item, _ in doc.iterate_items(root=table) if item is not table}


def _docx_docling_items(doc):
    """docling Word 문서 → _DocxItem 흐름.

    - 헤딩 레벨은 docling section_header.level 그대로(Heading 스타일 레벨과 같다). 'title' 라벨은 본문 줄이다.
    - 표 칸 안의 요소는 건너뛴다 — docling은 서식 있는 칸의 내용을 표와 별도 요소로 한 번 더 내놓는다. 칸 내용·칸 안 그림·
      중첩 표는 표 행(_docx_table_rows)에 이미 들어 있다.
    표는 TableBlock으로 내놓는다. 표와 그림은 from_table=True — 그림 캡션이 '제1장 …'로 시작해도 폴백 헤딩이 되지 않게.
    """
    from docling_core.types.doc import PictureItem, SectionHeaderItem, TableItem  # 지연 import

    inside_tables = _table_descendants(doc)
    for item, _ in doc.iterate_items():
        if item.self_ref in inside_tables:
            continue
        if isinstance(item, SectionHeaderItem):
            if text := _docling_norm(item.text):
                yield _DocxItem(True, item.level, text, False)
        elif isinstance(item, TableItem):
            if rows := _docx_table_rows(item, doc):
                yield _DocxItem(False, 0, TableBlock(rows), True)
        elif text := docling_block_text(item, doc):
            yield _DocxItem(False, 0, text, isinstance(item, PictureItem))


def _docx_sections(file_path: str | Path) -> list[Section]:
    """DOCX → docling → 섹션. Heading 스타일이 경계이고, 스타일이 하나도 없으면 조문 패턴 폴백(_docx_fallback_headings).

    헤딩은 본문에 넣지 않는다 — heading_path가 들고 있고 rag/index_text가 색인 입력 앞에 붙인다. 조립은 PDF와 같은
    sections.build_sections(본문 없는 헤딩 보존 포함). page는 없다(None).
    """
    items = list(_docx_docling_items(docling_convert(file_path)))
    return sections.build_sections(
        [(None, text, level if is_heading else None) for is_heading, level, text, _ in _docx_fallback_headings(items)]
    )


def extract_text(file_path: str | Path) -> str:
    """파일 전체를 텍스트 하나로 추출한다(청킹 없음). 채팅 첨부(컨텍스트 직접 주입)용.

    사용자 대기 경로라 docling을 안 태운다 — 표는 평문으로 흐르고, 이미지 속 글자(도표·스캔)는 읽지 못한다.
    """
    low = str(file_path).lower()
    if low.endswith('.pdf'):
        return '\n'.join(t for _, t in plain.pdf_pages(file_path)).strip()
    if low.endswith('.docx'):
        return _docx_text(file_path).strip()
    if low.endswith('.xlsx'):
        from rag.xlsx_chunking import (
            xlsx_plain_text,
        )  # 첨부는 openpyxl — docling 변환 프로세스를 색인과 나누지 않는다

        return xlsx_plain_text(str(file_path)).strip()  # 첨부는 행 묶음 없이 시트 전체
    return plain.read_text(file_path).strip()  # txt/md/기타


# MD 표 칸에서 docling이 망가뜨리는 두 표기 — 변환 전에 유니코드 사용자 영역 문자로 숨겼다가 결과에서 되돌린다.
# '\|'는 칸 구분으로 읽어 넘친 칸을 버리고, '<br>'은 그 행의 값을 표 밖 글로 내보낸다. '<br>' 변형은 '<br>' 하나로 되돌린다.
_MD_PIPE_MARK, _MD_BR_MARK = '\ue000', '\ue001'
_MD_BR_RE = re.compile(r'<br\s*/?>', re.IGNORECASE)


def _md_unmask(text: str) -> str:
    return text.replace(_MD_PIPE_MARK, '\\|').replace(_MD_BR_MARK, '<br>')


def _md_sections(file_path: str | Path) -> list[Section]:
    """마크다운 → docling(MD 백엔드, 같은 변환 프로세스) → 섹션. 표는 TableBlock으로, 헤딩은 문법대로(밑줄식 포함) 나온다.

    docling 출력을 그대로 쓰면 지는 곳 셋을 보정한다:
      ① 표 칸 안 '\\|'·'<br>' — 위 _MD_PIPE_MARK 주석
      ② 링크 — docling은 주소를 hyperlink 속성에만 둔다. '글(주소)'로 붙인다
      ③ 인라인 서식 — docling은 한 문장을 서식 단위 조각(같은 InlineGroup)으로 쪼갠다. 한 줄로 다시 잇고, 조각 사이
         띄어쓰기는 원문을 보고 정한다('**접수**로만' → '접수로만')
    레벨: '#'은 TitleItem(1단), '##'부터는 SectionHeaderItem(level 1부터) → +1. DOCX와 달리 MD의 '#'은 헤딩이다.
    빈 헤딩(YAML 앞머리)은 버리고, 코드 블록은 들여쓰기를 그대로 둔다. 글자 소실은 tests/test_chunking_md.py가 지킨다.

    알려진 한계(모양만 다르고 내용은 남는다):
      - ③ 원문에서 조각을 못 찾으면 띄운다 — 조사 앞에 빈칸이 생길 수 있다('확인 은')
      - 서식이 섞인 목록 항목은 번호가 따로 한 줄이 된다('1.' / '양식(…) 받기')
      - 그림('![대체 글](주소)')은 파일 안에 그림이 없어 캡션하지 않는다 — 대체 글과 자리표시가 남는다
      - 큰 표는 다른 형식과 같이 TABLE_MAX_TOKENS를 넘으면 색인하지 않는다
    """
    import tempfile

    from docling_core.types.doc import (
        CodeItem,
        InlineGroup,
        SectionHeaderItem,
        TableItem,
        TextItem,  # 지연 import
        TitleItem,
    )

    source = _MD_BR_RE.sub(_MD_BR_MARK, plain.read_text(file_path).replace('\\|', _MD_PIPE_MARK))
    with tempfile.NamedTemporaryFile('w', suffix='.md', encoding='utf-8', delete=False) as tmp:
        tmp.write(source)
    try:
        doc = docling_convert(tmp.name)
    finally:
        os.unlink(tmp.name)

    inside = _table_descendants(doc)
    items: list[Item] = []
    cursor = 0  # 원문에서 직전 글 조각이 끝난 자리 — ③ 띄어쓰기 판정
    group = None  # 직전 조각이 속한 InlineGroup
    for item, _ in doc.iterate_items():
        if item.self_ref in inside:
            continue
        parent = item.parent.resolve(doc) if item.parent else None
        inline = parent.self_ref if isinstance(parent, InlineGroup) else None
        raw = getattr(item, 'text', '') or ''
        found = source.find(raw, cursor) if raw else -1
        spaced = found < 0 or any(ch.isspace() for ch in source[cursor:found])
        if found >= 0:
            cursor = found + len(raw)
        if isinstance(item, (TitleItem, SectionHeaderItem)):
            if text := _md_unmask(_docling_norm(item.text)):
                items.append((None, text, 1 if isinstance(item, TitleItem) else item.level + 1))
        elif isinstance(item, TableItem):
            if rows := _docx_table_rows(item, doc):
                items.append((None, TableBlock([[_md_unmask(c) for c in row] for row in rows]), None))
        elif isinstance(item, CodeItem):
            # 코드 블록은 들여쓰기가 구조다 — 공백을 접는 _docling_norm을 태우지 않는다
            if text := _md_unmask('\n'.join(line.rstrip() for line in item.text.strip('\n').split('\n'))):
                items.append((None, text, None))
        else:
            text = docling_block_text(item, doc)
            if isinstance(item, TextItem) and item.hyperlink:
                text = f'{text}({item.hyperlink})'
            if text := _md_unmask(text):
                if inline and inline == group and items and items[-1][2] is None and isinstance(items[-1][1], str):
                    items[-1] = (None, items[-1][1] + (' ' if spaced else '') + text, None)
                else:
                    items.append((None, text, None))
        group = inline
    return sections.build_sections(items)


# 형식 → 섹션 파서. xlsx만 빠져 있다 — 행 묶음 + 시트 완성본·크기 초과 거절이라 계약이 다르다.
_SECTION_PARSERS = {
    '.pdf': _docling_sections,
    '.docx': _docx_sections,
    '.md': _md_sections,
    '.txt': plain.txt_sections,
}


def chunk_file(file_path: str | Path, *, description: str = '') -> list[ChunkData]:
    """파일 한 개를 청크 리스트로. 형식 분기는 여기 한 곳뿐이다 — 호출부는 확장자를 보지 않는다.

    xlsx는 chunk_xlsx로 위임한다(description은 xlsx 전용). 나머지는 섹션 추출 → sections.pack_sections → split.split_section.
    """
    suffix = Path(file_path).suffix.lower()
    splitter = SentenceSplitter(chunk_size=512, chunk_overlap=50)

    if suffix == '.xlsx':
        from rag.xlsx_chunking import chunk_xlsx  # 지연 import (xlsx_chunking → chunking 순환 방지)

        return chunk_xlsx(str(file_path), description=description)

    parser = _SECTION_PARSERS.get(suffix)
    if parser is None:
        # 업로드는 SUPPORTED_SUFFIXES가 앞에서 거르므로 도달하지 않는다. 도달했다면
        # 지원 안 되는 파일이 파서를 잘못 타고 있다는 뜻이라 조용히 넘기지 않는다.
        raise ValueError(f'지원하지 않는 형식: {suffix or "확장자 없음"}')

    out: list[ChunkData] = []
    table_keys: dict[tuple[int, int], str] = {}  # (섹션 번호, 섹션 안 블록 번호) → 문서 안 표 순번
    for si, section in enumerate(sections.pack_sections(parser(file_path))):
        for text, offset, table in split.split_section(section, splitter):
            # 이 조각에 실린 헤딩만 — 색인이 식별자 사전으로, 근거 준비가 구간 경계로 쓴다.
            lines = {line.strip() for line in text.split('\n')}
            heads = [h for h in section.headings if h.strip() in lines]
            meta = {'headings': heads} if heads else {}
            if table is not None:
                key = table_keys.setdefault((si, table.block), f't{len(table_keys) + 1}')
                meta |= {'is_table': True, 'table_key': key, 'table_text': table.text}
            out.append(
                ChunkData(
                    text=text,
                    heading_path=list(section.heading_path),
                    page=sections.page_at(section, offset),
                    chunk_index=len(out),
                    meta=meta or None,
                )
            )
    return out
