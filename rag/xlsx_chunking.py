"""xlsx를 시트별 행 묶음 청크로 나누고, 업로드 사전검증(표 설명 필수·표 크기 상한)을 한다.

- 청크: 시트 표를 약 `XLSX_BATCH_TOKENS` 크기의 행 묶음으로 나누고 묶음마다 표 설명·헤더를 반복한다 — 작은 묶음으로
  정밀하게 검색되게. 묶음마다 시트 완성본(`meta['table_text']`)을 실어, 근거 준비(`rag.retriever.prepare_evidence`)가
  걸린 묶음을 완성본으로 바꿔 모델이 언제나 시트 전체를 보게 한다. meta 계약은 PDF·DOCX 큰 표와 같다.
- 크기 상한: 시트 표 하나당 `TABLE_MAX_TOKENS` 생성 모델 토큰 — 행 수가 아니라 내용의 양으로 잰다(행당 토큰이 표 모양마다
  크게 달라서). 토큰은 렌더한 글자 수로 보수적으로 추정해(추정 ≥ 실제) 업로드가 GPU 서버 토크나이저에 묶이지 않는다.
  거절 문구는 토큰이라는 말 대신 허용 크기 대비 %와 행×열을 쓰고, 시트·파일로 나눠 올리라고 안내하지 않는다 — 나눈 조각은
  따로 검색돼 한 질문에 함께 걸린다는 보장이 없다. FE 업로드 도움말의 '셀 글자 수 합계 약 6천 자'는 이 상한을
  보수 계수로 환산한 값이다 — 상한·계수를 바꾸면 도움말도 바꾼다.
- 읽기: 색인 내용은 docling(안내문·양식이 섞인 시트를 표 단위로 더 잘 나눈다), 크기 판정은 openpyxl 렌더 — 업로드 검증이
  docling 변환을 기다리지 않게. 채팅 첨부(`xlsx_plain_text`)는 openpyxl로 읽고 묶지 않는다.
  docling은 완전히 빈 행·열을 표 경계로 본다 — 표 중간 빈 행에서 표가 갈려도 값은 같은 시트 청크에 이어진다.

상한 위반은 두 시점에 같은 함수로 걸린다: 웹 업로드는 라우터가 blob을 쓰기 전에 `validate_xlsx_upload`로 400을 주고,
라우터를 안 거치는 경로(재등재·재색인)는 `chunk_xlsx`의 예외로 워커가 문서를 failed로 굳힌다(재시도 안 하는 결정적 실패).
"""

import logging
import math
from contextlib import closing
from io import BytesIO

import openpyxl

# 표 markdown 조립은 chunking.py가 정의점 — 형식마다 표 모양이 다르면 프롬프트·리랭커·어휘 채널이 다른 모양을 본다.
from rag.chunking import (
    TABLE_BATCH_TOKENS,
    TABLE_MAX_TOKENS,
    ChunkData,
    docling_block_text,
    docling_convert,
    docling_grid_rows,
)
from rag.chunking import (
    batch_table_rows as _batch_rows,
)
from rag.chunking import (
    to_markdown_table as _to_markdown,
)
from rag.tokens import CONSERVATIVE_CHARS_PER_TOKEN, estimate_tokens_conservative

logger = logging.getLogger(__name__)

# 글자 → 토큰 추정 계수 — 정의점은 rag.tokens(근거 예산과 같은 추정기). 보수적이라 보통 표의 실제 상한은 명목보다 작다.
CHARS_PER_TOKEN = CONSERVATIVE_CHARS_PER_TOKEN
# 행 묶음 하나의 크기 — 정의점은 rag/chunking(PDF·DOCX 표도 같은 묶음 크기). 측정 근거 없이 잡은 값이다.
XLSX_BATCH_TOKENS = TABLE_BATCH_TOKENS


class XlsxUploadRejected(Exception):
    """xlsx 업로드 사전검증 실패의 공통 베이스 — 라우터는 이것 하나만 잡아 400으로 바꾼다.

    `str(e)`가 그대로 사용자에게 나간다 — 하위 클래스의 메시지는 사용자가 읽는 문장이다.
    """


class XlsxTableTooLarge(XlsxUploadRejected):
    """시트 표가 크기 상한을 넘었다. 문구는 토큰을 말하지 않고, 나눠 올리라고 하지 않는다(모듈 docstring)."""

    def __init__(self, sheet: str, percent: int, rows: int, cols: int):
        self.sheet, self.percent, self.rows, self.cols = sheet, percent, rows, cols
        super().__init__(
            f"시트 '{sheet}'의 표가 허용 크기의 {percent}%입니다(현재 {rows}행 × {cols}열). "
            "상담에 필요 없는 열이나 긴 설명을 줄여 다시 올려 주세요."
        )


class XlsxDescriptionRequired(XlsxUploadRejected):
    """xlsx인데 표 설명이 최종적으로 비었다.

    이번 요청에 안 보냈어도 직전 버전에서 계승되면 통과한다 — 계승 판정은 라우터가 끝낸 뒤 최종값만 넘긴다.
    """

    def __init__(self):
        super().__init__('xlsx 업로드에는 표 설명이 필요합니다 (검색 정확도를 위해 필수).')


def _cell(value) -> str:
    return '' if value is None else str(value)


def _has_value(value) -> bool:
    """공백만 든 칸은 빈칸이다 — docling과 기준이 다르면 openpyxl만 센 시트를 docling이 비워 내 색인이 실패한다."""
    return value is not None and str(value).strip() != ''


def description_missing(description: str | None) -> bool:
    """ "표 설명이 비었는가"의 정의점 — 업로드 사전검증과 PATCH 거절이 같은 판정을 쓴다.

    공백만도 "없음"이다 — 청크 조립이 `strip()`으로 설명 병합 여부를 가르므로 공백만은 검색에 기여하지 않는다.
    """
    return not (description or '').strip()


def estimate_tokens(text: str) -> int:
    """렌더한 표의 생성 모델 토큰 보수 추정(추정 ≥ 실제)."""
    return estimate_tokens_conservative(text)


def _check_sheet_size(title: str, header: list, body: list) -> None:
    """크기 상한 비교식 — 색인(`chunk_xlsx`)과 사전검증(`validate_xlsx_upload`)이 공유한다(openpyxl 렌더 기준)."""
    tokens = estimate_tokens(_to_markdown(header, body))
    if tokens > TABLE_MAX_TOKENS:
        raise XlsxTableTooLarge(title, math.ceil(tokens * 100 / TABLE_MAX_TOKENS), len(body), len(header))


def _iter_sheets(wb):
    """시트별 (시트명, 헤더행, 데이터행들) — 완전 빈 행은 제외한다.

    업로드 검증·색인·첨부가 모두 이 함수로 시트를 읽는다 — 재는 규칙이 두 벌이면 "업로드 통과·색인 거절"이 난다.
    숨긴 시트(hidden·veryHidden)는 검색 잡음이 되지 않게 읽지 않는다(크기 상한에도 안 센다). docling도 숨긴 시트를
    기본 순회에서 빼므로 두 경로의 시트 목록이 맞는다.
    """
    for ws in wb.worksheets:
        if ws.sheet_state != 'visible':
            continue
        rows = [
            [_cell(c.value) for c in row]
            for row in ws.iter_rows()
            if any(_has_value(c.value) for c in row)  # 완전 빈 행 스킵 — 공백만 든 칸도 빈칸
        ]
        if rows:
            yield ws.title, rows[0], rows[1:]


def validate_xlsx_upload(content: bytes, description: str | None) -> None:
    """xlsx 업로드 사전검증의 정의점 — 라우터가 blob을 쓰기 전에 부른다. 위반이면 `XlsxUploadRejected` 하위 예외.

    HTTP 변환은 라우터 몫이다(`rag/`는 웹 프레임워크를 들이지 않는다). 설명 검사부터 해 비었으면 워크북을 열지 않는다.
    docling 변환은 업로드 응답을 늦추므로 `chunk_xlsx`를 부르지 않고 openpyxl 렌더로만 잰다(색인과 같은 `_iter_sheets` 기준).
    """
    if description_missing(description):
        raise XlsxDescriptionRequired()
    with closing(openpyxl.load_workbook(BytesIO(content), read_only=True, data_only=True)) as wb:
        for title, header, body in _iter_sheets(wb):
            _check_sheet_size(title, header, body)


# 시트 블록: ('table', 행 목록[첫 행 = 헤더]) 또는 ('text', 본문) — 행 묶음을 행 단계에서 나누려고 표는 행 목록으로 든다.
Block = tuple[str, object]


def _render(block: Block) -> str:
    kind, value = block
    return _to_markdown(value[0], value[1:]) if kind == 'table' else value


def _sheet_chunks(title: str, blocks: list[Block], description: str, start: int) -> list[ChunkData]:
    """시트 하나 → 행 묶음 청크들. 표는 행 묶음으로 나누고, 표가 아닌 블록(그림 캡션·표 밖 문구)은 그대로 한 묶음.

    묶음마다 표 설명을 앞에 붙인다(작은 묶음일수록 검색 보강이 필요하다). meta['table_text']는 시트 완성본(설명 한 번 +
    블록들을 빈 줄로 이은 것) — 근거 준비 단계가 걸린 묶음을 '걸린 묶음 + 이것'으로 바꾼다.
    """
    prefix = f'[{description}]\n' if description.strip() else ''
    sheet_text = prefix + '\n\n'.join(_render(b) for b in blocks)
    if (est := estimate_tokens(sheet_text)) > TABLE_MAX_TOKENS:
        # 크기 판정 기준은 openpyxl 렌더다 — docling 완성본은 그림 캡션·병합 칸 반복·설명으로 더 클 수 있는데, 여기서
        # 거절하면 "업로드 통과·색인 거절"이 생기므로 경고만 남긴다. 근거 예산이 안전장치다.
        logger.warning(
            "시트 '%s' 완성본이 상한을 넘는다(추정 %d > %d토큰) — 색인은 계속한다", title, est, TABLE_MAX_TOKENS
        )
    budget = int(XLSX_BATCH_TOKENS * CHARS_PER_TOKEN)
    parts = [
        p for kind, value in blocks for p in (_batch_rows(value[0], value[1:], budget) if kind == 'table' else [value])
    ]
    return [
        ChunkData(
            text=prefix + p,
            heading_path=[title],
            page=None,
            chunk_index=start + k,
            meta={'is_table': True, 'table_key': title, 'table_text': sheet_text},
        )
        for k, p in enumerate(parts)
    ]


def _plain_sheets(file_path, *, check_size: bool = True) -> list[tuple[str, list[list[str]]]]:
    with closing(openpyxl.load_workbook(file_path, read_only=True, data_only=True)) as wb:
        out = []
        for title, header, body in _iter_sheets(wb):
            if check_size:
                _check_sheet_size(title, header, body)
            out.append((title, [header, *body]))
        return out


def xlsx_plain_text(file_path) -> str:
    """채팅 첨부용 — 시트별 표 markdown을 묶지 않고 그대로 잇는다.

    색인용 크기 상한은 걸지 않는다 — 첨부에는 더 좁은 자체 상한(schemas.kms.ATTACHMENT_MAX_TEXT_CHARS)이 있고, 여기서
    먼저 막으면 라우터의 일반 예외 처리에 걸려 거절 사유가 가려진다.
    """
    return '\n'.join(_render(('table', rows)) for _, rows in _plain_sheets(file_path, check_size=False))


def _check_workbook_size(file_path) -> list[str]:
    """크기 상한 검사(위반이면 XlsxTableTooLarge) 후 셀이 있는 시트 이름을 문서 순서로 돌려준다."""
    return [title for title, _ in _plain_sheets(file_path)]


def _docling_sheet_blocks(doc) -> dict[str, list[Block]]:
    """docling 엑셀 문서 → {시트명: [블록, …]}(읽기 순서). 시트는 docling의 'sheet' 그룹이고 그룹 이름이 시트명이다.

    한 시트에 표가 여럿이면 표마다 자기 헤더를 가진 블록이 된다(시트당 표 1개는 강제하지 않는다). 숨긴 시트의 그룹은
    기본 순회가 비워 내고, 블록이 없는 시트는 넣지 않는다.
    """
    from docling_core.types.doc import GroupLabel, TableItem  # 지연 import — docling은 톱레벨에 안 들인다

    sheets: dict[str, list[Block]] = {}
    for group in doc.groups:
        if group.label != GroupLabel.SHEET:
            continue
        blocks: list[Block] = []
        for item, _ in doc.iterate_items(root=group):
            if isinstance(item, TableItem):
                if rows := docling_grid_rows(item, doc):
                    blocks.append(('table', rows))
            elif text := docling_block_text(item, doc):
                blocks.append(('text', text))
        if blocks:
            sheets[group.name] = blocks
    return sheets


def chunk_xlsx(file_path, description: str = '') -> list[ChunkData]:
    """시트별 행 묶음 청크. description(표 설명)이 있으면 각 묶음 앞에 병합한다.

    크기 상한은 docling 변환 전에 openpyxl로 먼저 잰다(위반이면 XlsxTableTooLarge). openpyxl이 셀을 본 시트를 docling이
    비워 내면 조용히 빠뜨리지 않고 실패로 올린다. 셀 없이 그림만 있는 시트는 docling 결과가 있으면 뒤에 붙인다.
    """
    titles = _check_workbook_size(file_path)
    blocks = _docling_sheet_blocks(docling_convert(file_path))
    if missing := [t for t in titles if t not in blocks]:
        raise RuntimeError(f'docling이 셀이 있는 시트를 비워 냈다: {missing}')
    chunks: list[ChunkData] = []
    for t in titles + [t for t in blocks if t not in titles]:
        chunks += _sheet_chunks(t, blocks[t], description, len(chunks))
    return chunks
