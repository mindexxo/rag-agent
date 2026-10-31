"""#233 조문 인용 ≠ 조 제목 — 합성 문서 4건(재현 가능한 정본). 내용은 전부 지어낸 것이다.

실문서 전수(922줄)에는 인용을 제목으로 오인한 사례가 0건이라, 결함마다 경로를 골라 합성 문서로 고정한다(결정 6·설계 6):
  - aromanica_16 (DOCX, Heading 스타일 없음 → 조문 폴백): 인용 승격으로 문장 소실("제4조 및 제5조에 따라 …"·
                  "제6조 제2항의 …"·"제12조(이하 …)에 따른 …") / 본문 없는 '제7조(삭제)' 소실 / 괄호 중첩 제목 '제2조(시향 예약(단체))' /
                  '의' 생략 가지 조 '제 5 조1(분실 신고)'(#250)
  - summers_16  (DOCX, Heading 스타일): 본문 없는 조 — 중간('제3조(삭제)')·문서 끝('제6조(삭제)') 소실
  - homeplus_16 (MD): 본문 없는 조 '제3조(삭제)' 소실 / '(이하 …)' 인용 줄이 가짜 이름표 '9조'가 됨
  - goodpeople_16 (PDF, docling): 가짜 고아('제1장 총칙'이 제2장 첫 본문에 다시 붙음) / 문서 끝 '제7조(삭제)' 소실 /
                  '(이하 …)' 인용 줄이 가짜 이름표 '15조'가 됨 / 삭제 조가 다음 조 본문 앞에 붙어 "제5조" 질문에 제6조 내용이
                  넘어감(장을 넘는 '제2조(삭제)'·같은 장 '제5조(삭제)') — 이것은 #233에서 고치지 않은 알려진 한계(#260)다
골드 문항은 eval/gold_set_v2.jsonl의 *_ct0NN — 그중 삭제 조 배치 한계(#260)에 걸리는 5문항은 알려진 실패다(tests/test_citation_headings KNOWN_FAIL).
가짜 이름표 함정(문서에 없는 '제9조'·'제12조'·'제15조')은 골드가 아니라 테스트로만 본다(설계 11). 문서마다 '제N조'는 그 문서 안에 실제로 있는 조만 제목으로 쓰고, 이름표 함정의
'제9조'·'제12조'·'제15조'는 문서에 **없는** 조다(인용 문장에만 나온다 — 그 조의 이름표가 생기면 결함).

PDF는 build_table_docs.py와 같은 CID 한글 글꼴(HYGothic-Medium — 플랫폼 무관)을 쓴다. 장은 조보다 크게·앞 여백을 넓게,
본문 문단 사이에도 여백을 둔다 — 장을 조와 같은 크기로 두면 docling이 장을 본문으로 읽었고, 문단 여백이 없으면 이웃 문단을
한 줄로 합쳤다(2026-10-18 실측).

사용: python sample_docs/corpus_v2/_src/build_citation_docs.py   → sample_docs/corpus_v2/<tenant>/ 에 산출
"""
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[3]))

# 문서 작성 라이브러리(docx·reportlab·eval.table_parsing)는 build_* 함수 안에서 불러온다 — 테스트(tests/test_citation_headings)가
# 아래 문단 목록(원문 정본)만 쓰려고 이 모듈을 불러오므로, 톱레벨에 두면 그 라이브러리가 없을 때 순수 조립 테스트까지
# 수집 단계에서 죽는다(#233 리뷰 지적 2).
OUT = Path(__file__).resolve().parent.parent

# 문단 목록 — 각 문서의 원문 정본. 테스트(tests/test_citation_headings.py)가 이 목록으로 "원문 소실 0"을 검사한다.
AROMANICA = [
    '시향관 운영 지침 (ARO-OPS-016)',
    '제 1 장 총칙',
    '제1조(목적) 이 지침은 아로마니카 시향관의 예약과 시향 도구 대여·반납 기준을 정한다.',
    '제2조(시향 예약(단체))',
    '단체 시향은 6인 이상 12인 이하로 하며, 방문 4일 전까지 예약한다.',
    '제 2 장 도구 대여',
    '제3조(대여 도구) 시향관은 시향 스트립 묶음과 무향 커피빈 키트를 대여한다.',
    '제4조 및 제5조에 따라 대여한 시향 도구의 반납 기한은 대여일로부터 10일로 한다.',
    '제4조(반납 절차)',
    '반납은 시향관 안내 데스크에서 받으며, 반납 확인증을 발급한다.',
    '제12조(이하 "파손 보상 조항"이라 한다)에 따른 도구 파손 정산은 매주 목요일에 일괄 처리한다.',
    '제 5 조1(분실 신고)',
    '도구를 분실한 경우 24시간 안에 시향관에 신고한다.',
    '제6조(운영 시간)',
    '시향관은 화요일부터 일요일까지 오전 11시에서 오후 8시까지 운영한다.',
    '제6조 제2항의 휴관일은 매월 첫째 월요일로 한다.',
    '제7조(삭제)',
    '제8조(시행일) 이 지침은 2026년 11월 1일부터 시행한다.',
]

SUMMERS = [   # (Heading 레벨 또는 None, 글)
    (1, '매장 안전 관리 규정 (SUM-SAF-016)'),
    (2, '제1조(목적)'),
    (None, '이 규정은 썸머스 매장의 화재 예방과 비상 대피 기준을 정한다.'),
    (2, '제2조(소화기 점검)'),
    (None, '소화기는 매월 둘째 주 수요일에 압력계를 확인하고 점검표에 서명한다.'),
    (2, '제3조(삭제)'),
    (2, '제4조(비상 대피)'),
    (None, '비상 대피 훈련은 분기마다 1회 실시하며, 집결지는 매장 정문 앞 광장으로 한다.'),
    (2, '제5조(안전 교육)'),
    (None, '신규 직원은 첫 출근 주에 2시간의 안전 교육을 이수한다.'),
    (2, '제6조(삭제)'),
]

HOMEPLUS = """# PB 시식 행사 운영 기준 (HP-EVT-016)

## 제1조(목적)

이 기준은 홈플러스 PB 상품 시식 행사의 신청과 운영 절차를 정한다.

## 제2조(행사 신청)

시식 행사는 행사일 14일 전까지 점포 운영팀에 신청한다.

## 제3조(삭제)

## 제4조(시식 운영)

시식 행사 1회의 운영 시간은 3시간을 넘지 않는다.
제9조(이하 "위생 조항"이라 한다)에 따라 시식 도구는 행사 종료 즉시 세척하고 세척 일지에 기록한다.

## 제5조(시행일)

이 기준은 2026년 12월 1일부터 시행한다.
"""

GOODPEOPLE = [   # (스타일, 글) — 'h1' 문서 제목 / 'ch' 장 / 'h2' 조 제목 / 'p' 본문
    ('h1', '후원 물품 보관 규정 (GP-STO-016)'),
    ('ch', '제 1 장 총칙'),
    ('h2', '제 1 조(목적)'),
    ('p', '이 규정은 굿피플이 기부받은 후원 물품의 입고와 보관 기준을 정한다.'),
    ('h2', '제 2 조(삭제)'),
    ('ch', '제 2 장 보관'),
    ('h2', '제 3 조(보관 장소)'),
    ('p', '후원 물품은 물류센터 2층 보관실에 품목별로 구분하여 보관한다.'),
    ('h2', '제 4 조(보관 기간)'),
    ('p', '식품류 후원 물품의 보관 기간은 입고일로부터 90일을 넘지 않는다.'),
    ('p', '제15조(이하 "폐기 조항"이라 한다)에 따라 유통기한이 지난 물품은 매월 마지막 금요일에 폐기한다.'),
    ('h2', '제 5 조(삭제)'),
    ('h2', '제 6 조(재고 확인)'),
    ('p', '보관 물품 재고는 반기마다 전수 확인하고 결과를 이사회에 보고한다.'),
    ('h2', '제 7 조(삭제)'),
]


def build_aromanica() -> Path:
    """Heading 스타일이 하나도 없는 워드 — 전부 Normal 문단이라 조문 폴백(docx.docx_fallback_headings)을 탄다."""
    from docx import Document as DocxDocument
    p = OUT / 'aromanica' / 'aromanica_16_시향관운영지침.docx'
    d = DocxDocument()
    for text in AROMANICA:
        d.add_paragraph(text)
    d.save(p)
    return p


def build_summers() -> Path:
    from docx import Document as DocxDocument
    p = OUT / 'summers' / 'summers_16_매장안전관리규정.docx'
    d = DocxDocument()
    for level, text in SUMMERS:
        if level:
            d.add_heading(text, level=level)
        else:
            d.add_paragraph(text)
    d.save(p)
    return p


def build_homeplus() -> Path:
    p = OUT / 'homeplus' / 'homeplus_16_PB시식행사운영기준.md'
    p.write_text(HOMEPLUS, encoding='utf-8')
    return p


def build_goodpeople() -> Path:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate

    from eval.table_parsing.build import _BODY       # CID 한글 글꼴(HYGothic-Medium) 등록도 여기서 된다
    h1 = ParagraphStyle('h1', fontName='HYGothic-Medium', fontSize=16, leading=22, spaceAfter=6)
    ch = ParagraphStyle('ch', fontName='HYGothic-Medium', fontSize=14, leading=20, spaceBefore=14, spaceAfter=6)
    h2 = ParagraphStyle('h2', fontName='HYGothic-Medium', fontSize=13, leading=18, spaceBefore=8, spaceAfter=4)
    # 본문 문단 사이 여백 — 없으면 docling이 이웃 문단을 한 줄로 합친다(인용 줄이 앞 문장 뒤에 붙어 줄 머리가 아니게 된다, 실측)
    body = ParagraphStyle('p', parent=_BODY, spaceBefore=6)
    p = OUT / 'goodpeople' / 'goodpeople_16_후원물품보관규정.pdf'
    style = {'h1': h1, 'ch': ch, 'h2': h2, 'p': body}
    story = [Paragraph(text, style[kind]) for kind, text in GOODPEOPLE]
    SimpleDocTemplate(str(p), pagesize=A4, leftMargin=15 * mm, rightMargin=15 * mm).build(story)
    return p


if __name__ == '__main__':
    for build in (build_aromanica, build_summers, build_homeplus, build_goodpeople):
        print(build())
