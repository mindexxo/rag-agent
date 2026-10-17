"""#252 표 렌더 보강 문서 생성 — 실문서에 사례가 없거나 1건뿐인 결정을 덮는 합성 문서 3건(재현 가능한 정본).

#252 실문서 diff(102문서)에서 근거가 모자랐던 결정만 겨냥한다:
  - goodpeople_15 (PDF): 결정 8 — 조 헤딩 바로 아래 표(실문서는 docx 3건뿐, PDF 0건) /
                          결정 E·1 — 쪽을 넘기는 긴 표(쪽마다 헤더 반복 → 한 표로 잇고 행 묶음마다 헤더)
  - summers_15 (DOCX):    결정 4 — 표 안의 표 / 결정 9 — 가로 병합 값 / 결정 C — 칸 안 줄바꿈
  - harim_15 (XLSX):      결정 F — 날짜 칸(00:00:00 꼬리, 실문서 1건뿐) / 결정 9 — 가로 병합된 날짜 / 결정 C
골드 문항은 eval/gold_set_v2.jsonl의 *_tb00N(type=single_fact).

표 그리기는 eval/table_parsing/build.py(합성 표 하네스 #246)의 함수를 그대로 쓴다 — 병합·중첩·줄바꿈을 경로마다 같은
방식으로 만드는 정의점이 그쪽이다. 소재는 각 테넌트의 기존 문서와 겹치지 않는다(harim 제주 배송·명절 마감은 기존 문서에
있어 피했다 — #253 교훈 "한 사실은 한 문서에만").

사용: python sample_docs/corpus_v2/_src/build_table_docs.py   → sample_docs/corpus_v2/<tenant>/ 에 산출
"""
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[3]))

import openpyxl                                                    # noqa: E402
from docx import Document as DocxDocument                          # noqa: E402
from reportlab.lib.pagesizes import A4                             # noqa: E402
from reportlab.lib.styles import ParagraphStyle                    # noqa: E402
from reportlab.lib.units import mm                                 # noqa: E402
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer  # noqa: E402

from eval.table_parsing.build import _BODY, _docx_fill, _rl_table, _xlsx_fill   # noqa: E402

OUT = Path(__file__).resolve().parent.parent
_H1 = ParagraphStyle('h1', fontName='HYGothic-Medium', fontSize=16, leading=22, spaceAfter=6)
_H2 = ParagraphStyle('h2', fontName='HYGothic-Medium', fontSize=13, leading=18, spaceBefore=8, spaceAfter=4)

REGIONS = ['서울', '경기', '강원', '충북', '충남', '전북', '전남', '경북', '경남', '제주']


def volunteer_rows() -> list[list]:
    """나눔센터 70곳 — 쪽을 두 번 넘긴다. 값은 번호로 결정적(지역 10개 순환, 인원 4~15명)."""
    rows = [['번호', '봉사처', '지역', '월 배정 인원']]
    for i in range(1, 71):
        rows.append([str(i), f'나눔센터-{i:03d}', REGIONS[(i * 7) % 10], f'{(i * 5) % 12 + 4}명'])
    return rows


def build_goodpeople() -> Path:
    p = OUT / 'goodpeople' / 'goodpeople_15_자원봉사활동규정.pdf'
    story = [
        Paragraph('자원봉사 활동 규정 (GP-VOL-015)', _H1),
        # '제 N 장'이 있어야 조가 3단 헤딩이 되어 같은 장 안에서 묶인다(_sections_from_docling_elements) — 묶여야 조 헤딩이
        # 본문 줄로 내려가 결정 8(헤딩 바로 아래 표)이 생긴다. 장이 없는 PDF는 조마다 섹션이라 이 경우가 원리적으로 없다.
        Paragraph('제 1 장 총칙', _H2),
        Paragraph('제 1 조(목적)', _H2),
        Paragraph('이 규정은 굿피플 자원봉사자의 활동 수당, 배정 및 보험에 관한 기준을 정한다.', _BODY),
        Paragraph('제 2 조(적용 범위)', _H2),
        Paragraph('이 규정은 국내외 현장에서 활동하는 등록 자원봉사자에게 적용한다. 단기 행사 봉사자는 별도 공지에 따른다.', _BODY),
        Paragraph('제 2 장 수당과 배정', _H2),
        Paragraph('제 3 조(활동 수당)', _H2),
        # 결정 8: 헤딩과 표 사이에 글이 없다
        _rl_table([['활동 구분', '국내 1일', '해외 1일'],
                   ['현장 구호', '25,000원', '48,000원'],
                   ['교육 봉사', '18,000원', '41,000원'],
                   ['행정 지원', '15,000원', '37,000원']]),
        # 짧은 조를 하나 둔다 — 묶음(_pack_sections)은 작은 조끼리만 되므로 제3조(표)가 제4조와 한 섹션이 되어야 결정 8이 생긴다
        Paragraph('제 4 조(수당 지급일)', _H2),
        Paragraph('활동 수당은 활동 종료 다음 달 10일에 등록 계좌로 지급한다.', _BODY),
        Paragraph('제 5 조(봉사처별 배정 인원)', _H2),
        Paragraph('봉사처별 월 배정 인원은 다음과 같다.', _BODY),
        Spacer(1, 4),
        _rl_table(volunteer_rows()),                                # 결정 E·1: 쪽마다 헤더 반복(repeatRows=1)
        Paragraph('제 3 장 보험', _H2),
        Paragraph('제 6 조(보험)', _H2),
        Paragraph('해외 활동 자원봉사자는 출국 7일 전까지 단체 상해보험에 가입하며, 보험료는 법인이 전액 부담한다.', _BODY),
    ]
    SimpleDocTemplate(str(p), pagesize=A4, leftMargin=15 * mm, rightMargin=15 * mm).build(story)
    return p


def build_summers() -> Path:
    p = OUT / 'summers' / 'summers_15_매장장비점검지침.docx'
    d = DocxDocument()
    d.add_heading('매장 장비 점검 지침 (SUM-OPS-015)', level=1)
    d.add_paragraph('직영·팝업 매장의 장비 점검 주기와 보증 기간을 정한다.')

    def table(rows):
        t = d.add_table(rows=len(rows), cols=max(len(r) for r in rows))
        t.style = 'Table Grid'
        _docx_fill(t, rows)

    d.add_heading('1. 장비별 점검 기준', level=2)               # 결정 4: 표 안의 표
    table([['장비', '점검 기준'],
           ['진열 냉장고', {'nested': [['주기', '담당'], ['주 1회', '매장 부점장']]}],
           ['각인 기계', {'nested': [['주기', '담당'], ['월 2회', '본사 기술팀']]}],
           ['포스 단말기', '매일 개점 전 · 점장']])
    d.add_heading('2. 장비 보증 기간', level=2)                  # 결정 9: 두 매장 유형에 걸친 가로 병합
    table([['장비', '직영 매장', '팝업 매장'],
           ['진열 냉장고', {'t': '3년', 'cs': 2}, None],
           ['각인 기계', '2년', '1년'],
           ['포스 단말기', {'t': '18개월', 'cs': 2}, None]])
    d.add_heading('3. 정기 점검 시간', level=2)                  # 결정 C: 칸 안 줄바꿈
    table([['구분', '점검 시간'],
           ['평일', '오전 7시\n~ 오전 9시'],
           ['주말', '오후 10시\n~ 오후 11시 30분']])
    d.save(p)
    return p


def build_harim() -> Path:
    p = OUT / 'harim' / 'harim_15_협력농가위생점검일정표.xlsx'
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = '점검일정'
    _xlsx_fill(ws, [['농가 코드', '농가명', '1차 점검일', '2차 점검일', '비고'],
                    ['HF-101', '정읍 산들농장', '2026-11-04', '2027-01-13', '사료 교체\n후 재점검'],   # 결정 C
                    ['HF-102', '익산 햇살농장', '2026-11-06', '2027-01-15', ''],                      # 결정 F
                    ['HF-103', '김제 들녘농장', {'t': '2026-11-12', 'cs': 2}, None, '본사 합동 점검으로 1회 대체'],  # 결정 9
                    ['HF-104', '부안 바람농장', '2026-11-19', '2027-01-27', '']],
               typed=True)                                          # 날짜는 날짜 타입 칸(docling이 00:00:00을 붙이는 경로)
    wb.save(p)
    return p


if __name__ == '__main__':
    for build in (build_goodpeople, build_summers, build_harim):
        print(build())
