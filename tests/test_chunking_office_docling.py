"""DOCX·XLSX docling 경로 (#220) — 표 그리기 규칙은 가짜 grid로, 나머지는 생성한 파일을 실제로 변환해 고정한다.

Word·Excel은 SimplePipeline이라 레이아웃 모델이 없어 변환이 빠르다(파일당 1초 안팎). 그림 캡션은 VLM을 죽인 상태로
돌린다 — 외부 서비스 없이 결정적으로 돌고, '캡션 실패 = 자리표시' 경로를 함께 고정한다.
"""
from io import BytesIO
from types import SimpleNamespace

import openpyxl
import pytest

from config import settings
from rag.chunking import (_PICTURE_PLACEHOLDER, _docx_sections, chunk_file, count_picture_placeholders,
                          docling_grid_markdown)

pytest.importorskip('docling')

def _png(size=(120, 80)) -> bytes:
    # 1×1은 안 된다 — docling Word 백엔드가 아주 작은 그림을 여백(ContentLayer.INVISIBLE)으로 분류해 iterate_items에서 빠진다
    from PIL import Image
    buf = BytesIO()
    Image.new('RGB', size, (200, 30, 30)).save(buf, format='PNG')
    return buf.getvalue()


@pytest.fixture(autouse=True)
def _dead_vlm(monkeypatch):
    """VLM을 죽인 컨버터로 — 설정 변경이 먹으려면 싱글턴을 비워야 한다(test_chunking_docling과 같은 이유)."""
    import rag.chunking as ch
    monkeypatch.setattr(settings, 'docling_enabled', True)
    monkeypatch.setattr(settings, 'vlm_caption_url', 'http://127.0.0.1:1/v1/chat/completions')
    monkeypatch.setattr(settings, 'vlm_caption_timeout_seconds', 2.0)
    monkeypatch.setattr(ch, '_docling_converter', None)
    monkeypatch.setattr(ch, '_docling_semaphore', None)
    yield
    ch._docling_converter = None        # 죽은 URL이 박힌 컨버터를 다음 테스트에 남기지 않는다
    ch._docling_semaphore = None


def _both(path, monkeypatch):
    """같은 파일을 docling 경로·기존 경로로 각각 청킹."""
    new = chunk_file(path)
    monkeypatch.setattr(settings, 'docling_enabled', False)
    old = chunk_file(path)
    monkeypatch.setattr(settings, 'docling_enabled', True)
    return new, old


# ── 표 그리기 (docling_grid_markdown) — 가짜 grid ─────────────────────────────────

def _cell(text, r, c):
    return SimpleNamespace(text=text, start_row_offset_idx=r, start_col_offset_idx=c)


def _table(grid):
    return SimpleNamespace(data=SimpleNamespace(grid=grid))


class TestGridMarkdown:
    def test_가로_병합은_첫_열에만(self):
        title = _cell('업무 인수 인계서', 0, 0)
        md = docling_grid_markdown(_table([[title, title, title],
                                           [_cell('a', 1, 0), _cell('b', 1, 1), _cell('c', 1, 2)]]))
        assert md.splitlines()[0] == '| 업무 인수 인계서 |  |  |'
        assert md.count('업무 인수 인계서') == 1

    def test_세로_병합은_행마다_반복한다(self):
        # #220 결정 10 — 첫 행에만 두면 표가 청크 경계에서 잘릴 때 아래 행이 분류어를 잃는다(실문서 37행 중 7행)
        cat = _cell('발생원인', 1, 0)
        md = docling_grid_markdown(_table([[_cell('구분', 0, 0), _cell('유형', 0, 1)],
                                           [cat, _cell('기반구조 장애', 1, 1)],
                                           [cat, _cell('기술적 장애', 2, 1)]]))
        assert md.splitlines()[2:] == ['| 발생원인 | 기반구조 장애 |', '| 발생원인 | 기술적 장애 |']

    def test_값이_하나도_없는_행은_빠진다(self):
        md = docling_grid_markdown(_table([[_cell('h', 0, 0)], [_cell('', 1, 0)], [_cell('v', 2, 0)]]))
        assert md.splitlines() == ['| h |', '| --- |', '| v |']

    def test_빈_표는_None(self):
        assert docling_grid_markdown(_table([])) is None
        assert docling_grid_markdown(_table([[_cell('', 0, 0)]])) is None


# ── DOCX ─────────────────────────────────────────────────────────────────────────

def _docx(tmp_path, build, name='t.docx'):
    from docx import Document
    doc = Document()
    build(doc)
    p = tmp_path / name
    doc.save(p)
    return p


class TestDocxDocling:
    def test_헤딩_스타일_레벨이_기존_경로와_같다(self, tmp_path, monkeypatch):
        # Heading 스타일 문서 29건 전부 docling 레벨 = 스타일 레벨이었다(2026-10-11). PDF 매핑의 제목/장 휴리스틱을
        # 쓰면 3단이 2단으로 눌린다 — 이 테스트가 그 회귀를 잡는다.
        def build(doc):
            doc.add_heading('배송 정책', level=1)
            doc.add_heading('1. 배송 유형', level=2)
            doc.add_heading('1.1 일반 배송', level=3)
            doc.add_paragraph('일반 배송은 결제 후 2영업일 이내 출고한다.')
            doc.add_heading('1.2 새벽 배송', level=3)
            doc.add_paragraph('새벽 배송은 전날 23시까지 주문분이다.')
        p = _docx(tmp_path, build)
        new = [s.heading_path for s in _docx_sections(p)]     # 섹션 단위로 본다 — 청크는 형제 섹션을 묶어 공통 조상만 남긴다
        monkeypatch.setattr(settings, 'docling_enabled', False)
        old = [s.heading_path for s in _docx_sections(p)]
        assert new == old
        assert new == [['배송 정책', '1. 배송 유형', '1.1 일반 배송'], ['배송 정책', '1. 배송 유형', '1.2 새벽 배송']]

    def test_표_셀은_한_번만_들어가고_굵게_표시가_없다(self, tmp_path):
        # docling은 서식 있는 셀의 내용을 표와 별도 텍스트로 한 번 더 내놓는다 — 건너뛰지 않으면 셀이 두 번 색인된다.
        # export_to_dataframe은 굵은 셀을 '**…**'로 내보낸다 — 셀 원문으로 그려야 없다.
        # 중복은 **서식 있는 칸에서만** 생긴다 — 그래서 세는 문구를 굵은 칸에 둔다(평범한 칸으로 세면 건너뛰기를 꺼도 통과했다)
        def build(doc):
            doc.add_heading('요금표', level=1)
            t = doc.add_table(rows=2, cols=2)
            t.cell(0, 0).paragraphs[0].add_run('지역').bold = True
            t.cell(0, 1).text = '요금'
            t.cell(1, 0).paragraphs[0].add_run('제주고유지역').bold = True
            t.cell(1, 1).text = '5000원'
        chunks = chunk_file(_docx(tmp_path, build))
        text = '\n'.join(c.text for c in chunks)
        assert text.count('제주고유지역') == 1
        assert '**' not in text
        assert '| 지역 | 요금 |' in text

    def test_표_셀_안의_중첩_표를_읽는다(self, tmp_path, monkeypatch):
        # python-docx 경로는 셀 안의 표를 못 읽었다(인사규정 2건에서 문장 373·435개 누락)
        def build(doc):
            doc.add_heading('인사규정', level=1)
            outer = doc.add_table(rows=1, cols=1)
            inner = outer.cell(0, 0).add_table(rows=1, cols=2)
            inner.cell(0, 0).text = '중첩셀고유문구'
            inner.cell(0, 1).text = 'CL2'
        new, old = _both(_docx(tmp_path, build), monkeypatch)
        assert '중첩셀고유문구' in '\n'.join(c.text for c in new)
        assert '중첩셀고유문구' not in '\n'.join(c.text for c in old)

    def test_스타일_없는_문서는_조문_폴백이_docling_경로에서도_돈다(self, tmp_path):
        # #205 — Heading 스타일 0개인 실문서. docling도 헤딩을 0개 내놓으므로 폴백이 같은 층을 만들어야 한다.
        def build(doc):
            doc.add_paragraph('장애관리 지침서')
            doc.add_paragraph('제1장 총칙')
            doc.add_paragraph('제1조(목적) 이 지침은 장애 대응 절차를 정한다.')
            doc.add_paragraph('제2조(적용) 전 시스템에 적용한다.')
        assert [s.heading_path for s in _docx_sections(_docx(tmp_path, build))] == [
            ['장애관리 지침서', '제1장 총칙', '제1조(목적)'], ['장애관리 지침서', '제1장 총칙', '제2조(적용)']]

    def test_표_칸_안의_그림은_그_칸에_설명이_붙는다(self, tmp_path):
        # #220 리뷰 — 표 안 요소를 건너뛰는 규칙에 칸 안 그림까지 걸려 캡션이 버려졌다(자리표시도 없이). 이제 그 칸에 붙는다.
        # VLM이 죽었으니 자리표시가 칸에 들어가야 하고, 실패 지표도 그것을 센다.
        def build(doc):
            doc.add_heading('매뉴얼', level=1)
            t = doc.add_table(rows=1, cols=2)
            t.cell(0, 0).text = '1. 로그인 화면'
            t.cell(0, 1).paragraphs[0].add_run().add_picture(BytesIO(_png()))
        chunks = chunk_file(_docx(tmp_path, build))
        assert f'| 1. 로그인 화면 | {_PICTURE_PLACEHOLDER} |' in chunks[0].text
        assert count_picture_placeholders(chunks) == 1

    def test_빈_워드는_청크_0개(self, tmp_path):
        # 호출부(index_pending_document)가 '추출된 텍스트가 없습니다'로 failed 처리한다 — 여기선 예외 없이 빈 목록
        assert chunk_file(_docx(tmp_path, lambda doc: None)) == []

    def test_그림은_VLM이_죽으면_자리표시로_남는다(self, tmp_path):
        def build(doc):
            doc.add_paragraph('제1조(목적) 본문.')
            doc.add_picture(BytesIO(_png()))
            doc.add_paragraph('제2조(범위) 본문.')
        chunks = chunk_file(_docx(tmp_path, build))
        assert count_picture_placeholders(chunks) == 1

    def test_조문처럼_시작하는_그림_캡션은_헤딩이_되지_않는다(self, tmp_path, monkeypatch):
        # 그림은 from_table=True로 #205 폴백 승격에서 빠진다. 자리표시는 원래 조문 패턴에 안 걸리므로 캡션을 바꿔 끼워 본다
        import rag.chunking as ch
        monkeypatch.setattr(ch, 'docling_picture_text', lambda element: '제9조(그림 속 조문) 캡처 내용')

        def build(doc):
            doc.add_paragraph('제1조(목적) 본문.')
            doc.add_picture(BytesIO(_png()))
        sections = _docx_sections(_docx(tmp_path, build))
        assert all('제9조(그림 속 조문)' not in ' '.join(sec.heading_path) for sec in sections)
        assert any('제9조(그림 속 조문) 캡처 내용' in sec.body for sec in sections)


# ── XLSX ─────────────────────────────────────────────────────────────────────────

def _xlsx(tmp_path, build, name='t.xlsx'):
    wb = openpyxl.Workbook()
    build(wb)
    p = tmp_path / name
    wb.save(p)
    return p


class TestXlsxDocling:
    def test_가로_병합_제목은_한_번만_세로_병합_분류는_행마다(self, tmp_path):
        def build(wb):
            ws = wb.active
            ws.title = '양식'
            ws.append(['업무 교통비 청구서', None, None])
            ws.merge_cells('A1:C1')
            ws.append(['구분', '항목', '금액'])
            ws.append(['교통', '버스', 1500])
            ws.append([None, '택시', 29100])
            ws.merge_cells('A3:A4')
        chunks = chunk_file(_xlsx(tmp_path, build), description='교통비 양식')
        assert len(chunks) == 1
        text = chunks[0].text
        assert text.count('업무 교통비 청구서') == 1
        assert '| 교통 | 버스 | 1500 |' in text and '| 교통 | 택시 | 29100 |' in text
        assert text.startswith('[교통비 양식]\n')
        assert chunks[0].meta == {'is_table': True, 'sheet': '양식'}
        assert chunks[0].heading_path == ['양식']

    def test_시트_위_그림은_그_시트_청크에_들어간다(self, tmp_path):
        from openpyxl.drawing.image import Image as XlImage

        def build(wb):
            ws = wb.active
            ws.title = '첫시트'
            ws.append(['등급', '적립률'])
            ws.append(['VIP', '2.0'])
            ws.add_image(XlImage(BytesIO(_png())), 'D2')
            ws2 = wb.create_sheet('둘째')
            ws2.append(['항목', '값'])
            ws2.append(['x', '1'])
        chunks = chunk_file(_xlsx(tmp_path, build))
        assert [c.heading_path for c in chunks] == [['첫시트'], ['둘째']]     # 시트=1청크 유지
        assert _PICTURE_PLACEHOLDER in chunks[0].text and _PICTURE_PLACEHOLDER not in chunks[1].text

    def test_숨긴_시트는_색인하지_않고_행_상한에도_세지_않는다(self, tmp_path):
        # #220 리뷰 — 숨긴 시트가 있으면 문서 전체가 실패했다(openpyxl은 세고 docling은 비워 냄). 서비스 결정: 숨긴 시트는 건너뛴다
        from rag.xlsx_chunking import XLSX_MAX_ROWS, validate_xlsx_upload

        def build(wb):
            wb.active.title = '보임1'
            wb.active.append(['col'])
            wb.active.append(['v1'])
            hidden = wb.create_sheet('참조')
            hidden.append(['k'])
            for i in range(XLSX_MAX_ROWS + 5):      # 상한 초과 — 숨겼으니 거절 사유가 아니다
                hidden.append([i])
            hidden.sheet_state = 'hidden'
            very = wb.create_sheet('내부')
            very.append(['x'])
            very.append(['y'])
            very.sheet_state = 'veryHidden'
            wb.create_sheet('보임2').append(['col'])
            wb['보임2'].append(['v2'])
        p = _xlsx(tmp_path, build)
        assert [c.heading_path for c in chunk_file(p)] == [['보임1'], ['보임2']]
        validate_xlsx_upload(p.read_bytes(), '설명')        # 예외 없음

    def test_공백만_든_시트는_빈_시트로_건너뛴다(self, tmp_path):
        # docling은 공백 칸을 내용으로 보지 않는다 — openpyxl만 세면 '시트를 비워 냈다'로 문서 전체가 실패했다
        def build(wb):
            wb.active.append(['col'])
            wb.active.append(['v'])
            blank = wb.create_sheet('공백')
            blank['A1'] = ' '
            blank['B2'] = ''
        assert len(chunk_file(_xlsx(tmp_path, build))) == 1

    def test_셀_없이_그림만_있는_시트와_차트_시트는_보이는_시트_뒤에_붙는다(self, tmp_path):
        # openpyxl은 셀이 없는 시트·차트 시트를 세지 않는다(행 상한 대상 아님). docling이 내용(그림·차트)을 내면 뒤에 붙인다
        from openpyxl.chart import BarChart, Reference
        from openpyxl.drawing.image import Image as XlImage

        def build(wb):
            ws = wb.active
            ws.title = '데이터'
            for row in (['k', 'v'], ['a', 1], ['b', 2]):
                ws.append(row)
            wb.create_sheet('그림만').add_image(XlImage(BytesIO(_png())), 'B2')
            chart = BarChart()
            chart.add_data(Reference(ws, min_col=2, min_row=1, max_row=3), titles_from_data=True)
            wb.create_chartsheet('차트').add_chart(chart)
        chunks = chunk_file(_xlsx(tmp_path, build))
        assert chunks[0].heading_path == ['데이터']
        assert {c.heading_path[0] for c in chunks[1:]} == {'그림만', '차트'}
        assert all(c.meta == {'is_table': True, 'sheet': c.heading_path[0]} for c in chunks)

    def test_빈_엑셀은_청크_0개(self, tmp_path):
        assert chunk_file(_xlsx(tmp_path, lambda wb: None)) == []

    def test_비상_스위치를_끄면_openpyxl_경로와_같다(self, tmp_path, monkeypatch):
        from rag.xlsx_chunking import chunk_xlsx_plain

        def build(wb):
            wb.active.append(['col', 'v'])
            wb.active.append(['a', '1'])
        p = _xlsx(tmp_path, build)
        monkeypatch.setattr(settings, 'docling_enabled', False)
        assert [c.text for c in chunk_file(p)] == [c.text for c in chunk_xlsx_plain(p)]

    def test_docling이_셀_있는_시트를_비워_내면_실패한다(self, tmp_path, monkeypatch):
        # 조용히 시트를 빠뜨린 채 ready가 되지 않게 — PDF 경로의 불완전 결과 가드와 같은 규율
        import rag.xlsx_chunking as xc

        def build(wb):
            wb.active.append(['col'])
            wb.active.append(['1'])
        monkeypatch.setattr(xc, '_docling_sheet_texts', lambda doc: {})
        with pytest.raises(RuntimeError, match='비워 냈다'):
            chunk_file(_xlsx(tmp_path, build))

    def test_행_상한은_docling_변환_전에_openpyxl로_건다(self, tmp_path, monkeypatch):
        import rag.xlsx_chunking as xc
        from rag.xlsx_chunking import XLSX_MAX_ROWS, XlsxTooManyRows

        def build(wb):
            wb.active.append(['col'])
            for i in range(XLSX_MAX_ROWS + 1):
                wb.active.append([i])

        def boom(*a, **k):
            raise AssertionError('행 상한 위반인데 docling 변환까지 갔다')
        monkeypatch.setattr(xc, 'docling_convert', boom)
        with pytest.raises(XlsxTooManyRows):
            chunk_file(_xlsx(tmp_path, build))


# ── 채팅 첨부는 docling을 타지 않는다 (#220 결정 6) ──────────────────────────────────

class TestAttachmentStaysPlain:
    def test_첨부_추출은_docling_변환을_부르지_않는다(self, tmp_path, monkeypatch):
        # 첨부는 사용자 대기 경로 — docling 세마포어를 색인과 나누면 PDF 색인 뒤에 줄을 선다
        import rag.chunking as ch
        import rag.xlsx_chunking as xc
        from rag.chunking import extract_text

        def boom(*a, **k):
            raise AssertionError('첨부가 docling을 탔다')
        monkeypatch.setattr(ch, 'docling_convert', boom)
        monkeypatch.setattr(xc, 'docling_convert', boom)

        def build_x(wb):
            wb.active.append(['col'])
            wb.active.append(['첨부값'])

        def build_d(doc):
            doc.add_paragraph('첨부 본문')
        assert '첨부값' in extract_text(_xlsx(tmp_path, build_x))
        assert '첨부 본문' in extract_text(_docx(tmp_path, build_d))


# ── 변환기 설정 ────────────────────────────────────────────────────────────────────

class TestConverterOptions:
    def test_워드_엑셀_캡션_설정이_PDF와_같다(self):
        # 조용히 무시되는 설정을 잡는 테스트(#151 교훈) — 한 형식만 캡션이 빠지거나 다른 프롬프트로 돌면 에러 없이 품질만 갈린다
        from docling.datamodel.base_models import InputFormat
        from rag.chunking import _docling_runtime

        converter, _ = _docling_runtime()
        pdf = converter.format_to_options[InputFormat.PDF].pipeline_options
        for fmt in (InputFormat.DOCX, InputFormat.XLSX):
            opts = converter.format_to_options[fmt].pipeline_options
            assert opts.do_picture_description is True
            assert opts.enable_remote_services is True   # 빠지면 OperationNotAllowed로 변환 자체가 죽는다
            assert opts.picture_description_options == pdf.picture_description_options
