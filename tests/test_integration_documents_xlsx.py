"""문서 관리 통합 테스트 — xlsx 표 설명 필수·표 크기 사전검증(#222)·설정 계승 출처·설명 PATCH (#194).

공용 헬퍼는 tests/helpers_documents.py. 설정 계승 출처(C절)는 xlsx 전용이 아니지만 여기 둔다 —
"xlsx 설명을 계승하려면 출처가 ALIVE여야 한다"는 요구가 `latest_alive`로 기준을 좁힌 계기라
한 파일에서 읽히는 게 낫다.

E2E 목록(PR 본문)의 A·B·C·D·E·F1을 덮는다. F2(101쪽 PDF)·G(동시성·FE)는 수동으로 남긴다.
"""
import pytest
from sqlalchemy import delete as sql_delete, func, select

from database import AsyncSessionLocal
from rag import cache, outbox
from rag.models import AnswerCache as AnswerCacheRow, Document, SearchIndexOutbox
from schemas.kms import DOCUMENT_DESCRIPTION_MAX
from tests.conftest import ingest
from tests.helpers_documents import MD, chunk_texts, doc_data, get_doc, make_folder, upload
from tests.test_xlsx_chunking import _max_rows_at_limit, _xlsx_bytes

XLSX_MIME = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
TABLE = _xlsx_bytes(3)


async def _post_xlsx(client, filename: str, content: bytes = TABLE, description=None, form=None):
    """상태 코드를 직접 보는 xlsx 업로드. description은 document-data 파트로, form은 평면 Form으로."""
    files = {'file': (filename, content, XLSX_MIME)}
    if description is not None:
        files.update(doc_data(description=description))
    return await client.post('/kms/documents', files=files, data=form)


async def _upload_xlsx(client, filename: str, description: str, content: bytes = TABLE) -> dict:
    res = await _post_xlsx(client, filename, content, description=description)
    assert res.status_code == 200, res.text
    return res.json()


async def _pending_ops(tenant_id: str) -> list[str]:
    async with AsyncSessionLocal() as s:
        return list((await s.execute(
            select(SearchIndexOutbox.op).where(SearchIndexOutbox.tenant_id == tenant_id)
            .where(SearchIndexOutbox.status == 'pending'))).scalars().all())


async def _clear_outbox(tenant_id: str) -> None:
    async with AsyncSessionLocal() as s:
        await s.execute(sql_delete(SearchIndexOutbox).where(SearchIndexOutbox.tenant_id == tenant_id))
        await s.commit()


async def _doc_count(tenant_id: str) -> int:
    async with AsyncSessionLocal() as s:
        return (await s.execute(select(func.count()).select_from(Document)
                                .where(Document.tenant_id == tenant_id))).scalar()


def _blob_names(blob_tmp, tenant_id) -> list:
    d = blob_tmp / tenant_id
    return sorted(d.iterdir()) if d.exists() else []


async def _insert_failed(tenant_id: str, filename: str, *, version: int, blob_tmp, **cols) -> int:
    """failed 행을 직접 심는다 — 옛 코드가 남긴 이력, 또는 색인 장애로 굳은 문서를 재현."""
    async with AsyncSessionLocal() as s:
        row = Document(tenant_id=tenant_id, filename=filename, mime=XLSX_MIME,
                       blob_path=str(blob_tmp / tenant_id / f'failed_v{version}.xlsx'),
                       version=version, status='failed', status_reason='테스트가 심음', **cols)
        s.add(row)
        await s.commit()
        return row.id


# ── A. 업로드 — xlsx 표 설명 필수 ───────────────────────────────────────────

@pytest.mark.asyncio
async def test_A1_신규_xlsx_설명_없으면_400_흔적_없음(client, tenant_id, fake_queue, blob_tmp):
    """blob을 쓰기 전에 거절한다 — 디스크·DB·대기열 어디에도 흔적이 없어야 한다."""
    res = await _post_xlsx(client, '혜택표.xlsx')
    assert res.status_code == 400
    assert '표 설명' in res.json()['detail']
    assert _blob_names(blob_tmp, tenant_id) == []
    assert await _doc_count(tenant_id) == 0
    assert await _pending_ops(tenant_id) == []


@pytest.mark.asyncio
async def test_A2_설명_있으면_색인되고_청크_앞에_붙는다(client, tenant_id, fake_queue, blob_tmp):
    doc = await _upload_xlsx(client, '혜택표.xlsx', '멤버십 등급별 혜택')
    assert doc['status'] == 'pending' and doc['description'] == '멤버십 등급별 혜택'
    await ingest(doc['document_id'])
    assert (await get_doc(doc['document_id'])).status == 'ready'
    texts = await chunk_texts(doc['document_id'])
    assert len(texts) == 1 and texts[0].startswith('[멤버십 등급별 혜택]\n')


@pytest.mark.asyncio
async def test_A3_공백만은_미전송과_동일하게_400(client, tenant_id, fake_queue, blob_tmp):
    assert (await _post_xlsx(client, '혜택표.xlsx', description='  \t ')).status_code == 400


@pytest.mark.asyncio
async def test_A4_재업로드_설명_미전송이면_살아있는_직전_버전에서_계승(client, tenant_id, fake_queue, blob_tmp):
    v1 = await _upload_xlsx(client, '혜택표.xlsx', '원래 설명')
    await ingest(v1['document_id'])
    res = await _post_xlsx(client, '혜택표.xlsx')          # 설명 없이
    assert res.status_code == 200, res.text
    v2 = res.json()
    assert v2['version'] == 2 and v2['description'] == '원래 설명'


@pytest.mark.asyncio
async def test_A5_재업로드_설명_전송이면_덮어쓴다(client, tenant_id, fake_queue, blob_tmp):
    v1 = await _upload_xlsx(client, '혜택표.xlsx', '원래 설명')
    await ingest(v1['document_id'])
    v2 = await _upload_xlsx(client, '혜택표.xlsx', '새 설명')
    assert v2['version'] == 2 and v2['description'] == '새 설명'


@pytest.mark.asyncio
async def test_A6_failed만_있으면_설명_없는_재업로드는_400(client, tenant_id, fake_queue, blob_tmp):
    """failed는 "없는 문서"(#161) — 그 위에 올리는 건 신규 업로드라 설명이 필요하다(사용자 결정 2026-10-03).
    failed 행에 설명이 남아 있어도 계승 출처가 아니다. 설명을 보내면 그 행이 되살아난다(version 유지)."""
    await _insert_failed(tenant_id, '혜택표.xlsx', version=1, blob_tmp=blob_tmp, description='남은 설명')
    assert (await _post_xlsx(client, '혜택표.xlsx')).status_code == 400
    revived = await _upload_xlsx(client, '혜택표.xlsx', '다시 쓴 설명')
    assert revived['version'] == 1 and revived['description'] == '다시 쓴 설명'


@pytest.mark.asyncio
async def test_A7_전부_deleted면_설명_없는_재업로드는_400(client, tenant_id, fake_queue, blob_tmp):
    v1 = await _upload_xlsx(client, '혜택표.xlsx', '지워질 설명')
    await ingest(v1['document_id'])
    assert (await client.delete(f"/kms/documents/{v1['document_id']}")).status_code == 204
    assert (await _post_xlsx(client, '혜택표.xlsx')).status_code == 400


@pytest.mark.asyncio
async def test_A8_비xlsx는_설명_없어도_그대로_통과(client, tenant_id, fake_queue, blob_tmp):
    doc = await upload(client, '환불정책.md', MD)
    assert doc['status'] == 'pending' and doc['description'] is None


@pytest.mark.asyncio
async def test_A9_A10_201자_설명은_두_수신_경로_모두_422(client, tenant_id, fake_queue, blob_tmp):
    """평면 Form 경로는 #194 전까지 ValidationError를 안 잡아 500이 났을 자리다."""
    too_long = 'x' * (DOCUMENT_DESCRIPTION_MAX + 1)
    assert (await _post_xlsx(client, '혜택표.xlsx', description=too_long)).status_code == 422       # JSON 파트
    assert (await _post_xlsx(client, '혜택표.xlsx', form={'description': too_long})).status_code == 422  # 평면 Form
    assert _blob_names(blob_tmp, tenant_id) == []


# ── B. 업로드 — xlsx 표 크기 상한(#222, 옛 150행) ─────────────────────────────

@pytest.mark.asyncio
async def test_B1_크기_상한_초과는_업로드_즉시_400(client, tenant_id, fake_queue, blob_tmp):
    """전에는 200 pending → 1분 뒤 failed였다(#139 이후). 이제 라우터가 막는다. 문구는 #222 결정 9 그대로."""
    res = await _post_xlsx(client, '큰표.xlsx', _xlsx_bytes(_max_rows_at_limit() + 1), description='설명')
    assert res.status_code == 400
    detail = res.json()['detail']
    assert '허용 크기의' in detail and '%입니다(현재 ' in detail and '토큰' not in detail
    assert _blob_names(blob_tmp, tenant_id) == []
    assert await _doc_count(tenant_id) == 0


@pytest.mark.asyncio
async def test_B2_정확히_상한이면_통과(client, tenant_id, fake_queue, blob_tmp):
    doc = await _upload_xlsx(client, '경계표.xlsx', '설명', _xlsx_bytes(_max_rows_at_limit()))
    await ingest(doc['document_id'])
    assert (await get_doc(doc['document_id'])).status == 'ready'


@pytest.mark.asyncio
async def test_B3_손상된_xlsx는_400이지_500이_아니다(client, tenant_id, fake_queue, blob_tmp):
    res = await _post_xlsx(client, '깨진.xlsx', b'this is not a zip', description='설명')
    assert res.status_code == 400
    assert '열 수 없습니다' in res.json()['detail']
    assert _blob_names(blob_tmp, tenant_id) == []


# ── C. 설정 계승 출처 — ALIVE만 ──────────────────────────────────────────────

@pytest.mark.asyncio
async def test_C1_failed_버전이_더_높아도_살아있는_버전에서_계승(client, tenant_id, fake_queue, blob_tmp):
    """v1 ready(폴더 A) + v2 failed(폴더 B, 옛 코드 이력). 전에는 번호가 큰 v2에서 물려받았다 —
    "없는 문서"의 설정이 신규 업로드에 묻어가는 셈. 이제 v1에서 계승하고 번호만 v3다."""
    fa, fb = await make_folder(client, 'A'), await make_folder(client, 'B')
    v1 = await upload(client, '환불정책.md', MD, extra_parts=doc_data(folder_id=fa, description='X'))
    await ingest(v1['document_id'])
    await _insert_failed(tenant_id, '환불정책.md', version=2, blob_tmp=blob_tmp,
                         folder_id=fb, description='Y', is_searchable=False)

    v3 = await upload(client, '환불정책.md', MD)
    assert v3['version'] == 3
    assert v3['folder_id'] == fa and v3['description'] == 'X' and v3['is_searchable'] is True


@pytest.mark.asyncio
async def test_C2_전부_deleted면_아무것도_계승하지_않는다(client, tenant_id, fake_queue, blob_tmp):
    """exists·_current_version은 deleted를 "없음"으로 답한다 — 사용자에겐 신규 등록인데
    지운 문서의 폴더·설명이 묻어오면 안 된다."""
    fa = await make_folder(client, 'A')
    v1 = await upload(client, '환불정책.md', MD, extra_parts=doc_data(folder_id=fa, description='X'))
    await ingest(v1['document_id'])
    assert (await client.patch(f"/kms/documents/{v1['document_id']}",
                               json={'is_searchable': False})).status_code == 200
    assert (await client.delete(f"/kms/documents/{v1['document_id']}")).status_code == 204

    v2 = await upload(client, '환불정책.md', MD)
    assert v2['version'] == 2                                  # 번호는 deleted도 센다(UNIQUE)
    assert v2['folder_id'] is None and v2['description'] is None and v2['is_searchable'] is True


@pytest.mark.asyncio
async def test_C3_정상_계승은_그대로(client, tenant_id, fake_queue, blob_tmp):
    fa = await make_folder(client, 'A')
    v1 = await upload(client, '환불정책.md', MD, extra_parts=doc_data(folder_id=fa, description='X'))
    await ingest(v1['document_id'])
    v2 = await upload(client, '환불정책.md', MD)
    assert v2['folder_id'] == fa and v2['description'] == 'X'


# ── D. PATCH description ────────────────────────────────────────────────────

async def _patch(client, doc_id: int, **body):
    return await client.patch(f'/kms/documents/{doc_id}', json=body)


@pytest.mark.asyncio
async def test_D1_ready_xlsx_설명_변경은_재색인된다(client, tenant_id, fake_queue, blob_tmp):
    doc = await _upload_xlsx(client, '혜택표.xlsx', 'A')
    await ingest(doc['document_id'])
    await _clear_outbox(tenant_id)

    res = await _patch(client, doc['document_id'], description='B')
    assert res.status_code == 200, res.text
    assert res.json()['status'] == 'pending' and res.json()['description'] == 'B'
    assert await _pending_ops(tenant_id) == [outbox.INDEX_DOCUMENT]

    await ingest(doc['document_id'])
    assert (await get_doc(doc['document_id'])).status == 'ready'
    assert (await chunk_texts(doc['document_id']))[0].startswith('[B]\n')


@pytest.mark.asyncio
async def test_D2_같은_값_재전송은_재색인하지_않는다(client, tenant_id, fake_queue, blob_tmp):
    doc = await _upload_xlsx(client, '혜택표.xlsx', 'A')
    await ingest(doc['document_id'])
    await _clear_outbox(tenant_id)
    res = await _patch(client, doc['document_id'], description='A')
    assert res.status_code == 200 and res.json()['status'] == 'ready'
    assert await _pending_ops(tenant_id) == [outbox.META_DOCUMENTS]


@pytest.mark.asyncio
async def test_D3_xlsx_설명을_비우면_400(client, tenant_id, fake_queue, blob_tmp):
    doc = await _upload_xlsx(client, '혜택표.xlsx', 'A')
    for empty in ('', '   '):
        assert (await _patch(client, doc['document_id'], description=empty)).status_code == 400
    assert (await get_doc(doc['document_id'])).description == 'A'


@pytest.mark.asyncio
async def test_D4_비xlsx_설명_변경은_컬럼만_바꾼다(client, tenant_id, fake_queue, blob_tmp):
    doc = await upload(client, '환불정책.md', MD)
    await ingest(doc['document_id'])
    await _clear_outbox(tenant_id)
    res = await _patch(client, doc['document_id'], description='메모')
    assert res.status_code == 200 and res.json()['status'] == 'ready'
    assert (await get_doc(doc['document_id'])).description == '메모'
    assert await _pending_ops(tenant_id) == [outbox.META_DOCUMENTS]


@pytest.mark.asyncio
async def test_D5_failed_xlsx_설명_변경은_되살리지_않는다(client, tenant_id, fake_queue, blob_tmp):
    """되살리기는 재업로드 경로의 역할(#161). 설명만 고쳤는데 재처리가 돌면 안 된다."""
    did = await _insert_failed(tenant_id, '혜택표.xlsx', version=1, blob_tmp=blob_tmp, description='A')
    res = await _patch(client, did, description='B')
    assert res.status_code == 200 and res.json()['status'] == 'failed'
    assert (await get_doc(did)).description == 'B'
    assert await _pending_ops(tenant_id) == [outbox.META_DOCUMENTS]


@pytest.mark.asyncio
async def test_D6_pending_xlsx_설명_변경은_행을_더_만들지_않는다(client, tenant_id, fake_queue, blob_tmp):
    """대기 중인 INDEX_DOCUMENT 행이 처리될 때 최신 설명을 읽는다 — 두 번째 행은 중복이다."""
    doc = await _upload_xlsx(client, '혜택표.xlsx', 'A')         # ingest 안 함 → pending
    res = await _patch(client, doc['document_id'], description='B')
    assert res.status_code == 200 and res.json()['status'] == 'pending'
    ops = await _pending_ops(tenant_id)
    assert ops.count(outbox.INDEX_DOCUMENT) == 1
    await ingest(doc['document_id'])
    assert (await chunk_texts(doc['document_id']))[0].startswith('[B]\n')


@pytest.mark.asyncio
async def test_D7_xlsx_설명_변경은_캐시를_무효화한다(client, tenant_id, fake_queue, blob_tmp):
    """visibility 전이(on→off)만 보던 기존 조건은 설명 변경을 못 잡는다 — 내용이 바뀌면 근거가 낡는다."""
    doc = await _upload_xlsx(client, '혜택표.xlsx', 'A')
    await ingest(doc['document_id'])
    async with AsyncSessionLocal() as s:
        await cache.save_answer(s, tenant_id, '혜택이 뭐예요', '답', [], [doc['document_id']])
        await s.commit()
    assert (await _patch(client, doc['document_id'], description='B')).status_code == 200
    async with AsyncSessionLocal() as s:
        remain = (await s.execute(select(func.count()).select_from(AnswerCacheRow)
                                  .where(AnswerCacheRow.tenant_id == tenant_id))).scalar()
    assert remain == 0


@pytest.mark.asyncio
async def test_D8_D9_미전송은_무변화_201자는_422(client, tenant_id, fake_queue, blob_tmp):
    doc = await _upload_xlsx(client, '혜택표.xlsx', 'A')
    await ingest(doc['document_id'])
    assert (await _patch(client, doc['document_id'], is_searchable=False)).status_code == 200
    assert (await get_doc(doc['document_id'])).description == 'A'
    assert (await _patch(client, doc['document_id'],
                         description='x' * (DOCUMENT_DESCRIPTION_MAX + 1))).status_code == 422


# ── E. 응답 필드 ────────────────────────────────────────────────────────────

@pytest.mark.asyncio
async def test_E1_E2_description_indexed_at이_모든_응답에_나온다(client, tenant_id, fake_queue, blob_tmp):
    doc = await _upload_xlsx(client, '혜택표.xlsx', 'A')
    assert doc['description'] == 'A' and doc['indexed_at'] is None       # 업로드 직후
    await ingest(doc['document_id'])

    single = (await client.get(f"/kms/documents/{doc['document_id']}")).json()
    assert single['indexed_at'] is not None and single['indexed_at'].endswith('+09:00')   # KST(#181)
    listed = (await client.get('/kms/documents')).json()['items'][0]
    patched = (await _patch(client, doc['document_id'], is_searchable=False)).json()
    for body in (single, listed, patched):
        assert body['description'] == 'A' and body['indexed_at'] == single['indexed_at']


# ── F. PDF 쪽수 — 설정 불변식 ────────────────────────────────────────────────

def test_F1_쪽수_가드가_타임아웃보다_먼저_걸린다():
    """정상 최악 = 100쪽 ≈ 610초(쪽당 6.1초, #168 실측) + 그림 100장 ≈ 200초(동시 1, #223 실측) < 변환 마감
    ≤ 회차 예산(arq job_timeout) − 꼬리. 순서가 깨지면 상한 안의 정상 문서가 마감에 걸리거나(앞),
    첫 행조차 회차 예산에 못 들어가 arq 취소로 헛돈다(뒤)."""
    from config import settings
    from rag import outbox
    from rag.worker import WorkerSettings
    per_page_seconds, per_picture_seconds = 6.1, 2.0
    assert (settings.docling_max_num_pages * per_page_seconds + settings.docling_max_pictures * per_picture_seconds
            < settings.docling_document_timeout_seconds
            <= WorkerSettings.job_timeout - outbox.INDEX_TAIL_SECONDS)


# ── G. docling 경로 (#220) — 업로드 → 워커 색인까지 실제 경로 ─────────────────────

def _wb_bytes(build) -> bytes:
    from io import BytesIO
    import openpyxl
    wb = openpyxl.Workbook()
    build(wb)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


@pytest.mark.asyncio
async def test_G1_숨긴_시트는_업로드를_막지_않고_색인에서_빠진다(client, tenant_id, fake_queue, blob_tmp):
    """숨긴 시트가 있으면 문서 전체가 failed였다(리뷰). 서비스 결정: 숨긴 시트는 읽지 않는다 — 행 상한에도 안 센다."""
    def build(wb):
        wb.active.title = '보임'
        wb.active.append(['등급', '적립률'])
        wb.active.append(['VIP', '2.0'])
        hidden = wb.create_sheet('참조')
        hidden.append(['k'])
        for i in range(5000):                       # 크기 상한 초과 — 숨겼으니 거절 사유가 아니다
            hidden.append([i])
        hidden.sheet_state = 'hidden'
    doc = await _upload_xlsx(client, '숨김표.xlsx', '등급 적립률', _wb_bytes(build))
    await ingest(doc['document_id'])
    assert (await get_doc(doc['document_id'])).status == 'ready'
    texts = await chunk_texts(doc['document_id'])
    assert len(texts) == 1 and '| VIP | 2.0 |' in texts[0]


@pytest.mark.asyncio
async def test_G2_병합_제목은_한_번만_색인된다(client, tenant_id, fake_queue, blob_tmp):
    def build(wb):
        ws = wb.active
        ws.append(['업무 교통비 청구서', None, None])
        ws.merge_cells('A1:C1')
        ws.append(['구분', '항목', '금액'])
        ws.append(['교통', '버스', 1500])
    doc = await _upload_xlsx(client, '양식.xlsx', '교통비 양식', _wb_bytes(build))
    await ingest(doc['document_id'])
    assert (await get_doc(doc['document_id'])).status == 'ready'
    texts = await chunk_texts(doc['document_id'])
    assert '\n'.join(texts).count('업무 교통비 청구서') == 1     # 제목 블록·표 묶음으로 나뉘어도(#222) 제목은 한 번


@pytest.mark.asyncio
async def test_G4_큰_시트는_행_묶음으로_색인되고_근거는_시트_완성본이다(client, tenant_id, fake_queue, blob_tmp):
    """#222 — 색인(행 묶음·묶음마다 완성본) → 조회(fetch_chunk_map이 table_text를 돌려줌) → 근거 준비(완성본으로 확장)"""
    from rag.search import os_search
    from rag.search.retriever import prepare_evidence

    def build(wb):
        ws = wb.active
        ws.title = '단가'
        ws.append(['상품코드', '상품명', '단가'])
        for i in range(200):
            ws.append([f'A-{i:04d}', f'상품 {i}', 1000 + i])
    doc = await _upload_xlsx(client, '단가표.xlsx', '상품 단가표', _wb_bytes(build))
    await ingest(doc['document_id'])
    assert (await get_doc(doc['document_id'])).status == 'ready'
    resp = await os_search.client().search(index=os_search.settings.opensearch_index, body={
        'size': 100, '_source': ['chunk_id'], 'query': {'term': {'document_id': doc['document_id']}}})
    ids = sorted(int(h['_source']['chunk_id']) for h in resp['hits']['hits'])
    assert len(ids) > 1                                              # 행 묶음 여럿
    cmap = await os_search.fetch_chunk_map(ids)
    chunks = [cmap[i] for i in ids]
    assert all(c.is_table and c.table_text == chunks[0].table_text for c in chunks)
    ev = prepare_evidence(chunks[::-1], top_k=5)                     # 아무 묶음이 걸려도
    assert len(ev) == 1                                              # 한 칸으로 접히고
    whole = ev[0].text.split('[표 전체]\n', 1)[1]                      # 걸린 묶음 먼저 + 표 전체(#268 결정 4)
    assert ev[0].text.startswith('[검색에 걸린 행]\n')
    assert whole == chunks[0].table_text and whole.count('| 상품코드 |') == 1 and 'A-0000' in whole and 'A-0199' in whole   # 시트 전체


@pytest.mark.asyncio
async def test_G3_docx_중첩_표_내용이_색인된다(client, tenant_id, fake_queue, blob_tmp):
    """python-docx 경로는 셀 안의 표를 못 읽었다(인사규정 2건 문장 373·435개 누락) — docling 경로로 색인까지 확인."""
    from io import BytesIO
    from docx import Document as Docx
    d = Docx()
    d.add_heading('인사규정', level=1)
    d.add_paragraph('직급 체계는 아래와 같다.')
    inner = d.add_table(rows=1, cols=1).cell(0, 0).add_table(rows=1, cols=2)
    inner.cell(0, 0).text = '중첩표고유문구'
    inner.cell(0, 1).text = 'CL2'
    buf = BytesIO()
    d.save(buf)
    res = await client.post('/kms/documents', files={'file': ('규정.docx', buf.getvalue(),
        'application/vnd.openxmlformats-officedocument.wordprocessingml.document')})
    assert res.status_code == 200, res.text
    doc_id = res.json()['document_id']
    await ingest(doc_id)
    assert (await get_doc(doc_id)).status == 'ready'
    assert any('중첩표고유문구' in t for t in await chunk_texts(doc_id))


@pytest.mark.asyncio
async def test_G5_첨부_엑셀이_크면_추출_불가가_아니라_첨부_상한_안내가_나간다(client, tenant_id):
    """#222 리뷰 — 색인용 크기 상한을 첨부에도 걸었더니 라우터의 일반 예외 처리에 걸려 "추출할 수 없는 파일"(422)로
    사유가 가려졌다. 첨부는 자체 상한(6천 자)의 413 안내를 받아야 한다."""
    content = _xlsx_bytes(_max_rows_at_limit() + 50)
    res = await client.post('/kms/attachments/extract', files={'file': ('큰표.xlsx', content, XLSX_MIME)})
    assert res.status_code == 413, res.text
    assert '필요한 부분만 잘라 첨부' in res.json()['detail']


# ── H. 색인 시간 상한·캡션 실패 (#223) — 업로드 → 워커 drain까지 실제 경로 ─────────────────
# VLM은 로컬 대역(conftest.fake_vlm). 변환은 변환 프로세스에서, 분류·재시도는 outbox에서 — 둘의 이음매를 본다.

DOCX_MIME = 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'


def _docx_with_pictures(n: int) -> bytes:
    from io import BytesIO
    from docx import Document as Docx
    from PIL import Image
    png = BytesIO()
    Image.new('RGB', (120, 80), (200, 30, 30)).save(png, format='PNG')
    d = Docx()
    d.add_heading('화면 매뉴얼', level=1)
    d.add_paragraph('로그인 절차를 설명한다.')
    for _ in range(n):
        d.add_picture(BytesIO(png.getvalue()))
    buf = BytesIO()
    d.save(buf)
    return buf.getvalue()


async def _outbox_row(doc_id: int) -> SearchIndexOutbox:
    async with AsyncSessionLocal() as s:
        return (await s.execute(select(SearchIndexOutbox)
                                .where(SearchIndexOutbox.op == outbox.INDEX_DOCUMENT)
                                .where(SearchIndexOutbox.payload['document_id'].as_integer() == doc_id))).scalar_one()


@pytest.mark.asyncio
async def test_H1_VLM_장애면_문서는_대기로_재시도되고_복구되면_ready(client, tenant_id, fake_queue, blob_tmp, fake_vlm):
    """#223 결정 3 — 전에는 VLM이 죽어 있으면 그림이 자리표시로 남은 채 ready가 됐다(조용한 불완전 색인)."""
    fake_vlm.status = 500
    res = await client.post('/kms/documents', files={'file': ('매뉴얼.docx', _docx_with_pictures(1), DOCX_MIME)})
    assert res.status_code == 200, res.text
    doc_id = res.json()['document_id']
    await ingest(doc_id)
    row = await _outbox_row(doc_id)
    assert (await get_doc(doc_id)).status == 'pending'                     # failed가 아니다 — 일시 실패
    assert row.status == outbox.PENDING and row.attempts == 1 and '캡션' in row.last_error
    assert row.next_attempt_at is not None                                 # 백오프가 걸렸다

    fake_vlm.status = 200                                                  # VLM 복구
    await ingest(doc_id)                                                   # row_ids 경로는 백오프를 무시한다(테스트 계약)
    assert (await get_doc(doc_id)).status == 'ready'
    from tests.conftest import FAKE_CAPTION
    assert any(FAKE_CAPTION in t for t in await chunk_texts(doc_id))


@pytest.mark.asyncio
async def test_H2_그림이_상한을_넘으면_사유와_함께_바로_failed(client, tenant_id, fake_queue, blob_tmp, fake_vlm, monkeypatch):
    """#223 결정 4 — 결정적 실패라 재시도 없이 1회로 확정되고, 사유가 FE에 그대로 보인다."""
    from config import settings
    monkeypatch.setattr(settings, 'docling_max_pictures', 2)
    res = await client.post('/kms/documents', files={'file': ('캡처.docx', _docx_with_pictures(3), DOCX_MIME)})
    doc_id = res.json()['document_id']
    await ingest(doc_id)
    doc = await get_doc(doc_id)
    assert doc.status == 'failed' and '그림 3장' in doc.status_reason
    assert (await _outbox_row(doc_id)).attempts == 1
    assert fake_vlm.calls == 0                                             # 넘는 문서에 VLM 시간을 쓰지 않았다
