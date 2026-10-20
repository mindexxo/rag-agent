"""PDF·DOCX 큰 표 확장 통합 테스트(#257) — 업로드 → 색인(표 묶음 메타) → 조회 → 근거 준비(표 전체), 상한 초과 표의 색인 실패.

엑셀 쪽(시트 묶음 → 시트 완성본)은 엑셀 통합 테스트(test_integration_documents_xlsx) G4 항목이 덮는다. 여기는 같은 계약을 PDF·DOCX 경로가
지키는지 본다 — DOCX로 만든다(docling SimplePipeline이라 PDF보다 빠르고, 표 → _TableBlock 운반은 두 형식이 같은 길이다).
"""
from io import BytesIO

import pytest
from docx import Document as Docx

from rag import os_search
from tests.conftest import ingest
from tests.helpers_documents import get_doc

DOCX_MIME = 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'


def _docx(n_rows: int) -> bytes:
    """본문 + 큰 표(n_rows행) + 작은 표 — 실제 안내문처럼 표 앞뒤에 글이 있다."""
    d = Docx()
    d.add_heading('협력사 납품 기준', level=1)
    d.add_paragraph('이 기준은 협력사별 납품 단가와 납기를 정한다.')
    d.add_paragraph('협력사별 납품 단가는 아래와 같다.')
    t = d.add_table(rows=n_rows + 1, cols=3)
    for c, h in enumerate(['품목코드', '품목명', '단가']):
        t.cell(0, c).text = h
    for i in range(1, n_rows + 1):
        for c, v in enumerate([f'P-{i:04d}', f'납품 품목 {i}번', f'{1000 + i * 7:,}원']):
            t.cell(i, c).text = v
    d.add_paragraph('※ 단가는 부가세 별도이다.')
    d.add_paragraph('단가는 분기마다 다시 정한다.')
    small = d.add_table(rows=2, cols=2)
    for r, row in enumerate([('구분', '납기'), ('일반', '7일')]):
        for c, v in enumerate(row):
            small.cell(r, c).text = v
    buf = BytesIO()
    d.save(buf)
    return buf.getvalue()


async def _upload(client, name: str, content: bytes) -> int:
    res = await client.post('/kms/documents', files={'file': (name, content, DOCX_MIME)})
    assert res.status_code == 200, res.text
    return res.json()['document_id']


async def _indexed(doc_id: int):
    resp = await os_search.client().search(index=os_search.settings.opensearch_index, body={
        'size': 200, '_source': ['chunk_id'], 'query': {'term': {'document_id': doc_id}}})
    ids = sorted(int(h['_source']['chunk_id']) for h in resp['hits']['hits'])
    cmap = await os_search.fetch_chunk_map(ids)
    return [cmap[i] for i in ids]


@pytest.mark.asyncio
async def test_큰_표는_묶음마다_표_키와_표_전체가_색인되고_근거는_표_전체다(client, tenant_id, fake_queue, blob_tmp):
    from rag.retriever import prepare_evidence
    doc_id = await _upload(client, '납품기준.docx', _docx(150))
    await ingest(doc_id)
    assert (await get_doc(doc_id)).status == 'ready'
    chunks = await _indexed(doc_id)
    batches = [c for c in chunks if c.is_table]
    assert len(batches) >= 2                                              # 큰 표는 행 묶음 여럿
    assert {c.table_key for c in batches} == {'t1'} and len({c.table_text for c in batches}) == 1
    whole = batches[0].table_text
    assert whole.startswith('이 기준은 협력사별 납품 단가와 납기를 정한다.\n협력사별 납품 단가는 아래와 같다.\n| 품목코드 |')   # 표 앞 설명 여러 줄(#268)
    assert whole.endswith('| P-0150 | 납품 품목 150번 | 2,050원 |\n※ 단가는 부가세 별도이다.')   # 표 아래 주석(#268) — 표지 없는 다음 문장은 안 붙는다
    assert all(b.text.startswith('협력사별 납품 단가는 아래와 같다.\n| 품목코드 |') and '※' not in b.text for b in batches)   # 묶음 본문은 그대로
    assert whole.count('| 품목코드 |') == 1 and 'P-0001' in whole and 'P-0150' in whole
    small = [c for c in chunks if '| 구분 | 납기 |' in c.text]
    assert small and not any(c.is_table for c in small)                  # 작은 표는 표 묶음이 아니다
    ev = prepare_evidence(batches[::-1] + small, top_k=5)                  # 아무 묶음이 걸려도
    from rag.retriever import table_evidence_text
    assert [c.text for c in ev] == [table_evidence_text(batches[-1]), small[0].text]   # 한 칸으로 접혀 걸린 묶음 + 표 전체(#268)
    assert ev[0].text.endswith(f'[표 전체]\n{whole}')


@pytest.mark.asyncio
async def test_상한을_넘는_표가_있으면_색인_실패와_사유가_남는다(client, tenant_id, fake_queue, blob_tmp):
    doc_id = await _upload(client, '대형표.docx', _docx(400))
    result = await ingest(doc_id)
    assert result['failed'] == 1                                          # 결정적 실패 — 재시도 없이 한 번에
    doc = await get_doc(doc_id)
    assert doc.status == 'failed'
    assert "표('품목코드 | 품목명 | 단가')가 허용 크기의" in doc.status_reason and '(현재 400행 × 3열)' in doc.status_reason
    assert await _indexed(doc_id) == []                                   # 검색에 아무것도 남지 않는다
