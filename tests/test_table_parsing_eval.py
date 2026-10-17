"""#246 표 파싱 검증 하네스의 순수 로직 — docling·서버 없이. 기대 렌더 규칙(expected), markdown 표 파싱·채점(score).

채점기가 틀리면 측정 전체가 틀린다 — 조립이 틀린 채로 측정한 사고(AGENTS.md)를 막는 테스트다.
"""
from types import SimpleNamespace

from eval.table_parsing.expected import Cell, expected_cells, expected_rows, expected_tables, flat
flat_dims = flat
from eval.table_parsing.score import ActualTable, chunk_tables, docling_tables, match_tables, score_cells, score_rows
from eval.table_parsing.specs import BY_ID, SPECS, paths_of

S02 = BY_ID['S02_colspan']
S16 = BY_ID['S16_rowheader']


class TestSpecs:
    def test_정의는_23축_하나씩이고_격자가_반듯하다(self):
        axes = {}
        for s in SPECS:
            axes.setdefault(s['axis'], []).append(s['id'])
            for t in expected_tables(s):
                n, m = flat(t)
                for r, row in enumerate(t):
                    assert len(row) == m, (s['id'], r)      # 들쭉날쭉한 행은 격자 비교를 깨뜨린다
                for c in expected_cells(t):                  # 병합 범위가 격자 밖으로 안 나간다
                    assert c.row + c.row_span <= n and c.col + c.col_span <= m, (s['id'], c)
        assert all(len(v) == 1 for v in axes.values()), {k: v for k, v in axes.items() if len(v) != 1}
        assert len(axes) == 23
        assert all(set(paths_of(s)) <= {'img', 'pdf', 'docx', 'xlsx'} for s in SPECS)


class TestExpectedRows:
    def test_전_경로_공통_grid_규칙_병합은_가로_세로_모두_칸마다(self):
        # #252 결정 2: PDF도 docling_grid_rows — 2단 헤더를 점으로 잇지 않고 첫 행이 헤더, 둘째 헤더 행은 데이터 행
        for path in ('pdf', 'docx', 'xlsx', 'img'):
            rows = expected_rows(S02, S02['rows'], path)
            assert rows[0] == ['직급', '국내 출장', '국내 출장', '해외 출장', '해외 출장']   # 결정 9
            assert rows[1] == ['직급', '일비', '숙박비', '일비', '숙박비'] and len(rows) == 5

    def test_가로_머리글_표도_첫_행이_헤더다(self):
        # #252 전 PDF는 열 이름이 0,1,2…였다(결함 A). 이제 어느 경로든 첫 행
        rows = expected_rows(S16, S16['rows'], 'pdf')
        assert rows[0] == ['문서명', '휴양시설 이용 안내'] and rows[1] == ['개정일', '2026-03-01']

    def test_중첩_표는_L2에선_어느_경로든_평탄화_L1_docx만_바깥_안쪽_따로(self):
        s = BY_ID['V15_nested']
        for path in ('pdf', 'docx', 'img'):
            flat, = expected_tables(s, path)                      # L2: docling PDF 실측 모양(#253) 5×3 — DOCX도 #252 D로 같은 모양
            assert flat_dims(flat) == (5, 3)
        assert len(expected_tables(s, 'docx', 'l1')) == 2 and len(expected_tables(s, 'pdf', 'l1')) == 1
        cells = {(c.row, c.col): (c.text, c.row_span, c.col_span) for c in expected_cells(flat)}
        assert cells[(0, 1)] == ('세부', 1, 2) and cells[(1, 0)] == ('평가 항목', 3, 1)
        assert cells[(1, 1)] == ('항목', 1, 1) and cells[(3, 2)] == ('40', 1, 1) and cells[(4, 1)] == ('매년 12월', 1, 2)
        assert expected_rows(s, flat, 'docx')[1] == ['평가 항목', '항목', '배점']   # grid 규칙: 세로 병합 반복


class TestChunkTables:
    def _chunks(self, *texts):
        return [SimpleNamespace(text=t, meta=None) for t in texts]

    def test_구분선으로_표를_가르고_헤더를_잡는다(self):
        t = chunk_tables(self._chunks('머리말\n| a | b |\n| --- | --- |\n| 1 | 2 |\n본문\n| c |\n| --- |\n| 3 |'))
        assert t == [[['a', 'b'], ['1', '2']], [['c'], ['3']]]

    def test_청크_경계에서_갈린_긴_표는_앞_표에_잇는다(self):
        t = chunk_tables(self._chunks('| a | b |\n| --- | --- |\n| 1 | 2 |', '| 3 | 4 |\n| 5 | 6 |'))
        assert t == [[['a', 'b'], ['1', '2'], ['3', '4'], ['5', '6']]]

    def test_같은_헤더가_반복되는_행_묶음은_한_표다(self):
        # #252: 긴 표는 행 묶음마다 헤더를 반복한다 — 묶음 둘을 표 하나로 읽어야 셀 대조가 된다. 다른 헤더면 새 표
        t = chunk_tables(self._chunks('| a | b |\n| --- | --- |\n| 1 | 2 |', '| a | b |\n| --- | --- |\n| 3 | 4 |', '| c |\n| --- |\n| 9 |'))
        assert t == [[['a', 'b'], ['1', '2'], ['3', '4']], [['c'], ['9']]]

    def test_xlsx는_시트_완성본만_읽는다(self):
        c = [SimpleNamespace(text='[표]\n| a |\n| --- |\n| 1 |', meta={'sheet_text': '[표]\n| a |\n| --- |\n| 1 |\n| 2 |'}),
             SimpleNamespace(text='[표]\n| a |\n| --- |\n| 2 |', meta={'sheet_text': '[표]\n| a |\n| --- |\n| 1 |\n| 2 |'})]
        assert chunk_tables(c) == [[['a'], ['1'], ['2']]]


class TestScore:
    def _exp(self):
        return [Cell(0, 0, 1, 2, '제목'), Cell(1, 0, 1, 1, '값 1'), Cell(1, 1, 1, 1, '3,000원')]

    def test_전부_맞으면_오류_없음(self):
        act = ActualTable([Cell(0, 0, 1, 2, '제목'), Cell(1, 0, 1, 1, '값 1'), Cell(1, 1, 1, 1, '3,000원')], 2, 2)
        s = score_cells(self._exp(), act, (2, 2))
        assert s.cells_ok == 3 and (s.numeric_ok, s.numeric_total) == (2, 2) and s.span_ok == 1 and s.kinds == []

    def test_오독_누락_병합을_가른다(self):
        act = ActualTable([Cell(0, 0, 1, 1, '제목'), Cell(1, 0, 1, 1, ''), Cell(1, 1, 1, 1, '3,OOO원')], 2, 2)
        s = score_cells(self._exp(), act, (2, 2))
        assert s.kinds == ['누락', '병합', '오독'] and s.numeric_ok == 0
        assert s.presence_ok == 1                       # '제목'만 어딘가에 있다

    def test_행열이_다르면_구조이고_존재율은_따로_센다(self):
        act = ActualTable([Cell(0, 0, 1, 1, '제목'), Cell(0, 1, 1, 1, '값 1'), Cell(0, 2, 1, 1, '3,000원')], 1, 3)
        s = score_cells(self._exp(), act, (2, 2))
        assert not s.dims_ok and '구조' in s.kinds and s.presence_ok == 3

    def test_표를_못_찾으면_미검출(self):
        s = score_cells(self._exp(), None, (2, 2))
        assert s.kinds == ['미검출'] and not s.detected and s.cells_total == 3

    def test_공백_정규화는_하되_붙여쓰기는_오독이다(self):
        # OCR이 '지원 기준'을 '지원기준'으로 내는 것은 검색 토큰이 달라지므로 맞은 것으로 치지 않는다
        act = ActualTable([Cell(0, 0, 1, 2, '제목'), Cell(1, 0, 1, 1, '값  1'), Cell(1, 1, 1, 1, '3,000원')], 2, 2)
        assert score_cells(self._exp(), act, (2, 2)).cells_ok == 3
        act = ActualTable([Cell(0, 0, 1, 2, '제목'), Cell(1, 0, 1, 1, '값1'), Cell(1, 1, 1, 1, '3,000원')], 2, 2)
        assert score_cells(self._exp(), act, (2, 2)).kinds == ['오독']

    def test_표_짝짓기는_글자가_가장_많이_겹치는_것(self):
        a = ActualTable([Cell(0, 0, 1, 1, '국내'), Cell(1, 0, 1, 1, '임원')], 2, 1)
        b = ActualTable([Cell(0, 0, 1, 1, '해외'), Cell(1, 0, 1, 1, '임원')], 2, 1)
        exp = [[Cell(0, 0, 1, 1, '해외'), Cell(1, 0, 1, 1, '임원')], [Cell(0, 0, 1, 1, '국내'), Cell(1, 0, 1, 1, '임원')]]
        assert match_tables(exp, [a, b]) == [b, a]
        # 기대 표가 둘인데 둘 다 겹침이 없으면 짝을 안 짓는다(미검출). 하나뿐일 때의 예외는 아래 테스트
        assert match_tables([[Cell(0, 0, 1, 1, '없는글자')], [Cell(0, 0, 1, 1, '또없음')]], [a, b]) == [None, None]

    def test_L2_행_채점은_같은_규칙(self):
        s = score_rows([['a', 'b'], ['1', '2']], [['a', 'b'], ['1', '9']])
        assert s.cells_ok == 3 and s.kinds == ['오독'] and (s.numeric_ok, s.numeric_total) == (1, 2)


    def test_기대에_없는_병합이_생겨도_병합_오류다(self):
        act = ActualTable([Cell(0, 0, 1, 2, '제목'), Cell(1, 0, 1, 2, '값 1'), Cell(1, 1, 1, 1, '3,000원')], 2, 2)
        assert '병합' in score_cells(self._exp(), act, (2, 2)).kinds

    def test_존재율은_띄어쓰기를_무시한_정확_일치(self):
        # '5'가 '150,000' 안에 있다고 존재로 치면 안 된다(리뷰) — 부분 문자열이 아니라 칸 단위 일치
        exp = [Cell(0, 0, 1, 1, '5'), Cell(0, 1, 1, 1, '지원 기준')]
        act = ActualTable([Cell(0, 0, 1, 1, '150,000'), Cell(0, 1, 1, 1, '지원기준')], 1, 2)
        s = score_cells(exp, act, (1, 2))
        assert s.presence_ok == 1 and s.cells_ok == 0

    def test_존재율_풀을_주면_짝지은_표_밖도_센다(self):
        # 쪽 넘김으로 갈린 뒷 표의 글자를 "유실"이 아니라 "갈림"으로 가르기 위해서다
        exp = [Cell(0, 0, 1, 1, 'a'), Cell(1, 0, 1, 1, 'b')]
        act = ActualTable([Cell(0, 0, 1, 1, 'a')], 1, 1)
        assert score_cells(exp, act, (2, 1), {'a', 'b'}).presence_ok == 2

    def test_겹침이_없어도_기대_표가_하나면_남은_표와_짝짓는다(self):
        # OCR이 전부 틀린 표를 미검출로 세지 않는다 — 오독으로 집계된다
        a = ActualTable([Cell(0, 0, 1, 1, '틀린글자')], 1, 1)
        assert match_tables([[Cell(0, 0, 1, 1, '정답')]], [a]) == [a]


class TestDoclingDrift:
    """기대 렌더 규칙이 docling_core·rag/chunking의 실제 함수와 같은지 — 규칙을 손으로 베낀 것이라 드리프트를 여기서 잡는다."""

    def _item(self, spec, header_rows=1):
        from docling_core.types.doc import TableCell, TableData, TableItem
        cells = []
        for c in expected_cells(spec['rows']):
            cells.append(TableCell(start_row_offset_idx=c.row, end_row_offset_idx=c.row + c.row_span,
                                   start_col_offset_idx=c.col, end_col_offset_idx=c.col + c.col_span,
                                   row_span=c.row_span, col_span=c.col_span, text=c.text,
                                   column_header=c.row < header_rows))
        n, m = flat(spec['rows'])
        return TableItem(self_ref='#/tables/0', data=TableData(num_rows=n, num_cols=m, table_cells=cells))

    def test_기대_규칙은_rag_chunking의_grid_렌더와_같다(self):
        # #252부터 PDF·DOCX·XLSX 모두 docling_grid_rows 하나다 — 하네스 기대 규칙이 그 함수와 어긋나면 여기서 잡힌다
        from rag.chunking import docling_grid_rows
        for spec in (S02, BY_ID['S03_rowspan'], BY_ID['S04_form'], S16):
            for path in ('pdf', 'docx'):
                assert docling_grid_rows(self._item(spec)) == expected_rows(spec, spec['rows'], path), (spec['id'], path)

    def test_docling_tables는_병합_원점만_남기고_헤더행을_센다(self):
        from docling_core.types.doc import DoclingDocument
        doc = DoclingDocument(name='t')
        doc.add_table(data=self._item(S02, header_rows=2).data)     # add_table이 body 참조까지 건다
        t, = docling_tables(doc)
        assert (t.n_rows, t.n_cols, t.header_rows) == (5, 5, 2)
        assert len(t.cells) == len(expected_cells(S02['rows']))      # grid의 반복 칸이 아니라 원점만


class TestBuild:
    def test_정의_하나가_4경로_문서로_나오고_병합이_파일에_실린다(self, tmp_path):
        import openpyxl
        from docx import Document
        from eval.table_parsing.build import build_all
        made = build_all(tmp_path, {'S02_colspan'})
        assert sorted(p for _, p, _ in made) == ['docx', 'img', 'pdf', 'xlsx']
        assert all(f.exists() and f.stat().st_size > 0 for _, _, f in made)
        ws = openpyxl.load_workbook(tmp_path / 'S02_colspan.xlsx.xlsx').active
        assert {str(m) for m in ws.merged_cells.ranges} == {'A1:A2', 'B1:C1', 'D1:E1'}
        t = Document(tmp_path / 'S02_colspan.docx.docx').tables[0]
        assert [c.text for c in t.rows[0].cells] == ['직급', '국내 출장', '국내 출장', '해외 출장', '해외 출장']
        assert (tmp_path / 'png' / 'S02_colspan.png').exists()

    def test_paths_인자로_경로를_좁힌다(self, tmp_path):
        from eval.table_parsing.build import build_all
        assert [p for _, p, _ in build_all(tmp_path, {'S01_basic'}, {'xlsx'})] == ['xlsx']


class TestGallery:
    def test_샘플_갤러리는_축마다_카드_하나와_이미지를_만든다(self, tmp_path):
        from eval.table_parsing.gallery import build_gallery
        page = build_gallery(tmp_path)
        html = page.read_text()
        assert html.count('class="c"') == 23 and len(list((tmp_path / 'png').glob('*.png'))) == 23


class TestQuestions:
    def test_정의마다_사실_2개_경로마다_다른_질문이고_정답이_표의_칸_값이다(self):
        from eval.table_parsing.questions import ALL_QUESTIONS, FACTS
        from eval.table_parsing.expected import norm
        from eval.table_parsing.render_image import cell_text
        for s in SPECS:
            assert len(FACTS[s['id']]) == 2, s['id']
            cells = {norm(cell_text(c)) for t in expected_tables(s) for row in t for c in row if c is not None}
            texts = set()
            for p, qs in ALL_QUESTIONS[s['id']].items():
                for q in qs:
                    assert any(norm(q.expected) in c for c in cells), (s['id'], q.expected)   # 정답은 표 안 글자
                    texts.add(q.text)
            n_paths = len(ALL_QUESTIONS[s['id']])
            assert len(texts) == 2 * n_paths, (s['id'], '경로마다 질문이 달라야 한다')

    def test_판정은_띄어쓰기_쉼표_원_날짜표기를_무시한다(self):
        from eval.table_parsing.questions import judge
        assert judge('임원의 국내 출장 일비는 50,000원입니다.', '50,000')
        assert judge('지원기준은 재직 1년이상입니다', '재직 1년 이상')          # OCR 띄어쓰기 소실
        assert judge('₩3,250,000입니다', '₩3,250,000') and judge('3250000원', '₩3,250,000')
        assert judge('복귀일은 2026년 11월 5일입니다', '2026-11-05')
        assert not judge('해당 내용은 문서에서 확인할 수 없습니다.', '500,000원')


class TestE2EPure:
    def test_SSE_파서는_delta와_done을_뽑고_깨진_줄은_건너뛴다(self):
        from eval.table_parsing.e2e import _sse
        text = ('event: meta\ndata: {"conversation_id": 1}\n\n'
                'event: delta\ndata: {"text": "임원"}\n\n'
                ': ping\n\n'
                'event: delta\ndata: {"text": "은 50,000원"}\n\n'
                'event: done\ndata: {"finish_reason": "done", "citations": [{"filename": "a.pdf"}]}\n')
        ev = _sse(text)
        assert ''.join(d['text'] for e, d in ev if e == 'delta') == '임원은 50,000원'
        assert next(d for e, d in ev if e == 'done')['citations'][0]['filename'] == 'a.pdf'

    def test_리뷰_요약은_축_경로별_정답_인용_수를_센다(self, tmp_path):
        from eval.table_parsing.e2e import phase_report
        manifest = [dict(id='S02_colspan', path='pdf', filename='S02_colspan.pdf.pdf', document_id=1, status='ready',
                         chunks=[{'text': '| a |', 'heading_path': []}]),
                    dict(id='S02_colspan', path='docx', filename='x', document_id=2, status='failed', chunks=[])]
        answers = [dict(id='S02_colspan', path='pdf', document_id=1, question='q1', expected='50,000', manual=False,
                        answer='50,000원', cited=['S02_colspan.pdf.pdf'], cited_ok=True, value_ok=True),
                   dict(id='S02_colspan', path='pdf', document_id=1, question='q2', expected='40,000', manual=False,
                        answer='모름', cited=[], cited_ok=False, value_ok=False)]
        (tmp_path / 'docs' / 'png').mkdir(parents=True)
        phase_report(tmp_path, manifest, answers)
        summary = (tmp_path / 'summary.md').read_text()
        assert '색인 1/1 · 생성 1/2(기계 1) · 인용 1/2' in summary and '색인 0/1 · 생성 0/0' in summary
        assert (tmp_path / 'review.html').read_text().count('class="c"') == 2


class TestLlmJudge:
    def test_재판정은_기계_실패분만_묻고_결과를_행에_쓴다(self):
        import asyncio
        from eval.table_parsing import llm_judge
        from eval.table_parsing.llm_judge import AnswerJudgment, rejudge, final_ok

        class FakeLlm:
            calls = []
            async def acomplete(self, messages, extra_body=None):
                FakeLlm.calls.append(messages[1]['content'])
                return AnswerJudgment(label='correct' if '30일' in messages[1]['content'] else 'ungrounded', reason='이유').model_dump_json()
        rows = [dict(id='S03_rowspan', path='pdf', question='무급 병가 한도?', expected='연 30일', answer='1년에 30일까지', value_ok=False, manual=False),
                dict(id='S20_symbols', path='img', question='백업 완료?', expected='☑', answer='완료 상태입니다', value_ok=False, manual=True),
                dict(id='S01_basic', path='pdf', question='x', expected='500,000원', answer='500,000원', value_ok=True, manual=False)]
        chunks = {('S03_rowspan', 'pdf'): [{'text': '| 병가 | 무급 | 연 30일 |'}], ('S20_symbols', 'img'): [{'text': '| 백업확인 |  | 운영팀 |'}]}
        done = asyncio.run(rejudge(rows, chunks, llm=FakeLlm()))
        assert len(done) == 2 and len(FakeLlm.calls) == 2                      # 기계가 맞다고 한 행은 안 묻는다
        assert '| 병가 | 무급 | 연 30일 |' in FakeLlm.calls[0]                    # 청크를 같이 준다
        assert rows[0]['llm_label'] == 'correct' and final_ok(rows[0])
        assert rows[1]['llm_label'] == 'ungrounded' and not final_ok(rows[1])  # 근거 없는 정답은 최종 성공이 아니다
        assert final_ok(rows[2]) and 'llm_label' not in rows[2]
        assert rows[0]['judge'] == llm_judge.judge_label()

    def test_판정_실패는_기계_판정을_유지하고_오류를_남긴다(self):
        import asyncio
        from eval.table_parsing.llm_judge import rejudge, final_ok

        class BrokenLlm:
            async def acomplete(self, messages, extra_body=None):
                raise RuntimeError('한도')
        rows = [dict(id='a', path='pdf', question='q', expected='e', answer='x', value_ok=False, manual=False)]
        asyncio.run(rejudge(rows, {}, llm=BrokenLlm()))
        assert rows[0]['llm_label'] is None and '한도' in rows[0]['llm_error'] and not final_ok(rows[0])


class TestDiagnosis:
    def test_청킹_손실은_청크에_없는_칸_글자만(self):
        from eval.table_parsing.e2e import chunk_loss
        spec = BY_ID['V14_vertical']
        chunks = [{'text': '| 0 | 1 | 2 |\n| --- | --- | --- |\n|  |  |  |\n| 교통비 | 150,000 | 매월25일 |\n| 통신비 | 50,000 | 매월 25일 |\n| 식대 | 200,000 | 매월 25일 |'}]
        assert chunk_loss(spec, 'img', chunks) == ['정액 항목', '월 지급액', '지급일']       # 세로쓰기 머리글만 없음(띄어쓰기는 무시)

    def test_실패_원인은_인용·청크·답변_모양으로_가른다(self):
        from eval.table_parsing.e2e import diagnose
        ch = [{'text': '| 식대 | 200,000 | 매월 25일 |'}]
        base = dict(expected='매월 25일', cited=[], cited_ok=False, final_ok=False, llm_label=None)
        assert '후보 0건' in diagnose(dict(base, answer='해당 내용은 제공된 문서에서 확인할 수 없습니다.'), ch)
        assert '파싱 단계 손실' in diagnose(dict(base, answer='확인할 수 없습니다.'), [{'text': '| 식대 | 200,000 |  |'}])
        assert '추정' in diagnose(dict(base, answer='완료', llm_label='ungrounded'), ch)
        assert '파싱 손실' in diagnose(dict(base, answer='매월 30일', llm_label='wrong'), [{'text': '| 식대 | 200,000 | 매월 30일 |'}])
        assert '생성 모델 오류' in diagnose(dict(base, answer='매월 30일', cited=['x'], cited_ok=True, llm_label='wrong'), ch)
        assert diagnose(dict(base, final_ok=True, answer='매월 25일'), ch) == ''
