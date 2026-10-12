"""엑셀 크기 추정의 보수성 (#222 결정 8) — 글자 수 ÷ CHARS_PER_TOKEN 추정이 **실제 생성 모델 토큰 이상**인가.

업로드 검증은 GPU 서버에 묶이지 않으려고 글자 수로 추정한다. 추정이 실제보다 작으면 상한(6,000토큰)을 넘는 표가
통과해 품질 꺾임 구간(실측 1만 토큰 92%·표 중간 행 15/19)에 들어간다. 그래서 실제 토크나이저(vLLM /tokenize)로
대조한다 — 서버에 닿지 않으면 이 파일만 건너뛴다(이유를 남긴다). 표본: 모의 코퍼스 xlsx 전부 + 열 수·셀 길이를
바꾼 합성 표(실측에서 글자/토큰이 가장 낮았던 쪽 = 숫자 많은 표 포함).
"""
import random
from pathlib import Path

import httpx
import openpyxl
import pytest

from config import settings
from contextlib import closing

from rag.xlsx_chunking import _iter_sheets, _to_markdown, estimate_tokens

ROOT = Path(__file__).resolve().parent.parent
BASE = settings.vllm_base_url.rstrip('/').removesuffix('/v1')


def _count(text: str) -> int:
    r = httpx.post(f'{BASE}/tokenize', json={'model': settings.vllm_model, 'prompt': text}, timeout=20)
    r.raise_for_status()
    return r.json()['count']


@pytest.fixture(scope='module', autouse=True)
def _need_vllm():
    try:
        _count('ping')
    except Exception as e:                       # noqa: BLE001
        pytest.skip(f'vLLM /tokenize에 닿지 않는다 — 추정 보수성 검사를 건너뛴다: {type(e).__name__}')


def _synthetic(tmp_path, cols: int, long: bool, numeric: bool) -> Path:
    rnd = random.Random(cols * 10 + long)
    words = '등급 적립 배송 반품 교환 쿠폰 회원 결제 할인 기간 대상 제외 지역 상품 가격 수량 비고 코드 담당 처리'.split()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append([f'항목{i}' for i in range(cols)])
    for _ in range(80):
        ws.append([rnd.randint(1, 999999) if numeric or i % 3 == 0
                   else ' '.join(rnd.choice(words) for _ in range(6 if long else 1)) for i in range(cols)])
    p = tmp_path / f's_{cols}_{long}_{numeric}.xlsx'
    wb.save(p)
    return p


def _mixed(tmp_path) -> Path:
    """영문·URL·긴 숫자·날짜·코드가 섞인 표 — 리뷰 지적: 한글·짧은 숫자 표본만으로는 계수 여유(1.1 vs 실측 최소 1.142)가 얇다."""
    rnd = random.Random(7)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(['Order ID', 'URL', '계좌번호형 숫자', '날짜', 'SKU', 'Note'])
    for i in range(80):
        ws.append([f'ORD-{rnd.randint(10**9, 10**10)}', f'https://shop.example.com/p/{rnd.randint(1, 99999)}?ref=kms',
                   str(rnd.randint(10**13, 10**14)), f'2026-{rnd.randint(1, 12):02d}-{rnd.randint(1, 28):02d}',
                   f'SKU{rnd.randint(1000, 9999)}-X{i}', 'Express delivery, refund within 7 days'])
    p = tmp_path / 'mixed.xlsx'
    wb.save(p)
    return p


def _tables(tmp_path):
    files = sorted((ROOT / 'sample_docs').rglob('*.xlsx'))
    files += [_synthetic(tmp_path, c, lg, num) for c in (2, 5, 10, 15) for lg in (False, True) for num in (False, True)]
    files.append(_mixed(tmp_path))
    for f in files:                                       # 상한 검사 없이 운영과 같은 렌더만 — 여기선 추정치 비교가 목적
        with closing(openpyxl.load_workbook(f, read_only=True, data_only=True)) as wb:
            for title, header, body in _iter_sheets(wb):
                yield f'{f.name}:{title}', _to_markdown(header, body)


def test_추정_토큰은_실제_토큰_이상이다(tmp_path):
    under = [(name, estimate_tokens(md), _count(md)) for name, md in _tables(tmp_path)
             if estimate_tokens(md) < _count(md)]
    assert not under, f'추정이 실제보다 작은 표: {under[:5]}'
